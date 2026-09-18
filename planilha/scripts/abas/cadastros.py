"""Abas de cadastro: clientes, fornecedores, plano de contas, centros de custo,
contas bancárias e cartões."""
from openpyxl.chart import BarChart, Reference

from lib import dados_exemplo as D
from lib import estilos as E
from lib import parametros as P
from lib import plano_contas as PC
from lib import util as U
from lib.util import L, R, nome, put

L_ENT = P.rng("ENTIDADE")
L_TPARC = P.rng("TPARC")
L_TIPO = P.rng("TIPO")
L_COMP = P.rng("COMP")
L_VENC = P.rng("VENC")
L_VPREV = P.rng("VPREV")
L_VFINAL = P.rng("VFINAL")
L_SALDO = P.rng("SALDO")
L_SUBCAT = P.rng("SUBCAT")
L_CAT = P.rng("CAT")
L_CC = P.rng("CC")
L_CARTAO = P.rng("CARTAO")
L_BANCO = P.rng("BANCO")
L_PAGTO = P.rng("PAGTO")
L_VDRE = P.rng("VDRE")
L_LINHA = P.rng("LINHADRE")
L_STATUS = P.rng("STATUS")

HOJE = "TODAY()"


def _cabecalho(ws, titulo, subtitulo, headers, larguras, primeira, ultima, altura=32):
    U.larguras(ws, larguras)
    U.titulo(ws, 1, titulo, subtitulo, span=min(len(headers), 12))
    U.cabecalho(ws, 5, 1, headers, "TH", altura)
    ws.row_dimensions[4].height = 6
    return primeira, ultima


def _linhas_vazias_input(ws, colunas, primeira, ultima):
    for col in colunas:
        for r in range(primeira, ultima + 1):
            ws.cell(row=r, column=col).style = "IN_TXT"


# ==========================================================================
# CLIENTES
# ==========================================================================
def clientes(wb, ctx):
    ws = wb.create_sheet(P.ABA_CLI)
    headers = ["Código", "Razão social", "Nome fantasia", "CPF/CNPJ", "Telefone", "WhatsApp", "E-mail",
               "Cidade", "UF", "Segmento", "Vendedor", "Limite de crédito", "Data de cadastro", "Status",
               "Observações", "Total comprado", "Qtd. de compras", "Ticket médio", "Última compra",
               "Valor em aberto", "Valor vencido", "Dias em atraso (maior)", "Situação"]
    larg = {"A": 9, "B": 40, "C": 24, "D": 19, "E": 15, "F": 16, "G": 30, "H": 16, "I": 5, "J": 18,
            "K": 16, "L": 15, "M": 13, "N": 11, "O": 34, "P": 15, "Q": 11, "R": 14, "S": 13, "T": 15,
            "U": 14, "V": 12, "W": 18}
    _cabecalho(ws, "CADASTRO DE CLIENTES",
               "Preencha as colunas amarelas. As colunas cinzas (P a W) são calculadas automaticamente "
               "a partir da aba LANÇAMENTOS — mantenha a razão social idêntica à usada nos lançamentos.",
               headers, larg, P.CLI_INI, P.CLI_FIM)
    estilos_in = ["IN_TXT_C", "IN_TXT", "IN_TXT", "IN_TXT_C", "IN_TXT_C", "IN_TXT_C", "IN_TXT",
                  "IN_TXT", "IN_TXT_C", "IN_TXT", "IN_TXT", "IN_MONEY", "IN_DATE", "IN_TXT_C", "IN_TXT"]
    for i, est in enumerate(estilos_in):
        col = i + 1
        for r in range(P.CLI_INI, P.CLI_FIM + 1):
            ws.cell(row=r, column=col).style = est
    for r in range(P.CLI_INI, P.CLI_FIM + 1):
        for col in range(16, 24):
            ws.cell(row=r, column=col).style = "CALC_TXT_C" if col in (17, 22) else \
                "CALC_MONEY" if col in (16, 18, 20, 21) else "CALC_DATE" if col == 19 else "CALC_TXT"
    for i, c in enumerate(D.CLIENTES):
        r = P.CLI_INI + i
        valores = [c["cod"], c["razao"], c["fantasia"], c["doc"], c["tel"], c["whats"], c["email"],
                   c["cidade"], c["uf"], c["segmento"], c["vendedor"], c["limite"], c["cadastro"],
                   c["status"], c["obs"]]
        for j, v in enumerate(valores):
            ws.cell(row=r, column=j + 1, value=v)
    # indicadores automáticos
    for r in range(P.CLI_INI, P.CLI_FIM + 1):
        b, c = f"$B{r}", f"$C{r}"
        nome_expr = f'IF({c}={b},0,SUMIFS({L_VPREV},{L_ENT},{c}))'
        qtd_expr = f'IF({c}={b},0,COUNTIFS({L_ENT},{c},{L_TPARC},1))'
        put(ws, r, 16, f'=IF({b}="","",SUMIFS({L_VPREV},{L_ENT},{b})+{nome_expr})', "CALC_MONEY")
        put(ws, r, 17, f'=IF({b}="","",COUNTIFS({L_ENT},{b},{L_TPARC},1)+{qtd_expr})', "CALC_INT")
        put(ws, r, 18, f'=IFERROR(IF($P{r}=0,"",$P{r}/$Q{r}),"")', "CALC_MONEY")
        put(ws, r, 19, f'=IF({b}="","",IFERROR(IF(SUMPRODUCT(MAX(({L_ENT}={b})*{L_COMP}))>0,'
                      f'SUMPRODUCT(MAX(({L_ENT}={b})*{L_COMP})),""),""))', "CALC_DATE")
        put(ws, r, 20, f'=IF({b}="","",SUMIFS({L_SALDO},{L_ENT},{b},{L_TIPO},"Entrada")'
                       f'+IF({c}={b},0,SUMIFS({L_SALDO},{L_ENT},{c},{L_TIPO},"Entrada")))', "CALC_MONEY")
        put(ws, r, 21, f'=IF({b}="","",SUMIFS({L_SALDO},{L_ENT},{b},{L_TIPO},"Entrada",{L_VENC},"<"&{HOJE})'
                       f'+IF({c}={b},0,SUMIFS({L_SALDO},{L_ENT},{c},{L_TIPO},"Entrada",{L_VENC},"<"&{HOJE})))',
            "CALC_MONEY")
        put(ws, r, 22, f'=IF(OR({b}="",$U{r}=0),0,MAX(0,{HOJE}-IFERROR(-SUMPRODUCT(MAX(({L_ENT}={b})*'
                       f'({L_SALDO}>0)*({L_VENC}<{HOJE})*(-{L_VENC}))),{HOJE})))', "CALC_INT")
        put(ws, r, 23, f'=IF({b}="","",IF($N{r}="Bloqueado","⛔ Bloqueado",'
                       f'IF($U{r}>0,"⚠ Com valores vencidos","✔ Em dia")))', "CALC_TXT")
    # validações
    rng_status = U.faixa(14, P.CLI_INI, 14, P.CLI_FIM)
    U.dv_lista(ws, rng_status, "STATUS_CADASTRO", "Status inválido", "Escolha Ativo, Inativo ou Bloqueado.")
    U.dv_lista(ws, U.faixa(9, P.CLI_INI, 9, P.CLI_FIM), "UF", "UF inválida", "Escolha uma UF da lista.")
    U.dv_decimal(ws, U.faixa(12, P.CLI_INI, 12, P.CLI_FIM), "greaterThanOrEqual", "0",
                 "Limite inválido", "O limite de crédito deve ser maior ou igual a zero.")
    U.dv_data(ws, U.faixa(13, P.CLI_INI, 13, P.CLI_FIM))
    ws["B4"] = "⚠ DADOS DE EXEMPLO — os clientes abaixo são fictícios. Substitua pelos seus clientes reais."
    ws["B4"].style = "NOTA"
    U.cf_contem(ws, U.faixa(23, P.CLI_INI, 23, P.CLI_FIM), "Com valores vencidos", E.YEL_BG, E.YEL_TX)
    U.cf_contem(ws, U.faixa(23, P.CLI_INI, 23, P.CLI_FIM), "Bloqueado", E.RED_BG, E.RED_TX)
    U.cf_contem(ws, U.faixa(23, P.CLI_INI, 23, P.CLI_FIM), "Em dia", E.GREEN_BG, E.GREEN_TX)
    nome(wb, "CLIENTES_NOMES", f"'{P.ABA_CLI}'!$B${P.CLI_INI}:$B${P.CLI_FIM}")
    nome(wb, "CLIENTES_FANTASIA", f"'{P.ABA_CLI}'!$C${P.CLI_INI}:$C${P.CLI_FIM}")
    U.cf_maior(ws, U.faixa(21, P.CLI_INI, 21, P.CLI_FIM), 0, E.RED_BG, E.RED_TX)
    U.cf_barras(ws, U.faixa(16, P.CLI_INI, 16, P.CLI_FIM))
    U.preparar(ws, freeze="B6", filtro=f"A5:W{P.CLI_FIM}", repetir_linhas="5:5", tab_cor=E.BLUE)
    ctx.clientes = ws
    ctx.clientes_nomes = f"'{P.ABA_CLI}'!$C${P.CLI_INI}:$C${P.CLI_FIM}"
    return ws


# ==========================================================================
# FORNECEDORES
# ==========================================================================
def fornecedores(wb, ctx):
    ws = wb.create_sheet(P.ABA_FORN)
    headers = ["Código", "Razão social", "Nome fantasia", "CNPJ/CPF", "Contato", "Telefone", "E-mail",
               "Cidade", "UF", "Categoria principal", "Condição de pagamento", "Prazo médio (dias)",
               "Limite de crédito", "Status", "Observações", "Total comprado", "Total pago",
               "Total em aberto", "Valor vencido", "Última compra"]
    larg = {"A": 9, "B": 40, "C": 22, "D": 19, "E": 18, "F": 16, "G": 30, "H": 16, "I": 5, "J": 24,
            "K": 20, "L": 12, "M": 15, "N": 11, "O": 34, "P": 15, "Q": 15, "R": 15, "S": 14, "T": 13}
    _cabecalho(ws, "CADASTRO DE FORNECEDORES",
               "Preencha as colunas amarelas. As colunas cinzas (P a T) são calculadas automaticamente "
               "a partir da aba LANÇAMENTOS.", headers, larg, P.FORN_INI, P.FORN_FIM)
    estilos_in = ["IN_TXT_C", "IN_TXT", "IN_TXT", "IN_TXT_C", "IN_TXT", "IN_TXT_C", "IN_TXT", "IN_TXT",
                  "IN_TXT_C", "IN_TXT", "IN_TXT", "IN_INT", "IN_MONEY", "IN_TXT_C", "IN_TXT"]
    for i, est in enumerate(estilos_in):
        for r in range(P.FORN_INI, P.FORN_FIM + 1):
            ws.cell(row=r, column=i + 1).style = est
    for r in range(P.FORN_INI, P.FORN_FIM + 1):
        for col in (16, 17, 18, 19):
            ws.cell(row=r, column=col).style = "CALC_MONEY"
        ws.cell(row=r, column=20).style = "CALC_DATE"
    for i, f in enumerate(D.FORNECEDORES):
        r = P.FORN_INI + i
        valores = [f["cod"], f["razao"], f["fantasia"], f["doc"], f["contato"], f["tel"], f["email"],
                   f["cidade"], f["uf"], f["categoria"], f["condicao"], f["prazo"], f["limite"],
                   f["status"], f["obs"]]
        for j, v in enumerate(valores):
            ws.cell(row=r, column=j + 1, value=v)
    for r in range(P.FORN_INI, P.FORN_FIM + 1):
        b = f"$B{r}"
        put(ws, r, 16, f'=IF({b}="","",SUMIFS({L_VPREV},{L_ENT},{b},{L_TIPO},"Saída"))', "CALC_MONEY")
        put(ws, r, 17, f'=IF({b}="","",SUMIFS({L_VFINAL},{L_ENT},{b},{L_TIPO},"Saída"))', "CALC_MONEY")
        put(ws, r, 18, f'=IF({b}="","",SUMIFS({L_SALDO},{L_ENT},{b},{L_TIPO},"Saída"))', "CALC_MONEY")
        put(ws, r, 19, f'=IF({b}="","",SUMIFS({L_SALDO},{L_ENT},{b},{L_TIPO},"Saída",{L_VENC},"<"&{HOJE}))',
            "CALC_MONEY")
        put(ws, r, 20, f'=IF({b}="","",IFERROR(IF(SUMPRODUCT(MAX(({L_ENT}={b})*{L_COMP}))>0,'
                      f'SUMPRODUCT(MAX(({L_ENT}={b})*{L_COMP})),""),""))', "CALC_DATE")
    U.dv_lista(ws, U.faixa(14, P.FORN_INI, 14, P.FORN_FIM), "STATUS_CADASTRO", "Status inválido",
               "Escolha Ativo, Inativo ou Bloqueado.")
    U.dv_lista(ws, U.faixa(9, P.FORN_INI, 9, P.FORN_FIM), "UF", "UF inválida", "Escolha uma UF da lista.")
    U.dv_inteiro(ws, U.faixa(12, P.FORN_INI, 12, P.FORN_FIM), "between", "0", "365",
                 "Prazo inválido", "Informe o prazo médio em dias (0 a 365).")
    U.cf_maior(ws, U.faixa(19, P.FORN_INI, 19, P.FORN_FIM), 0, E.RED_BG, E.RED_TX)
    U.cf_barras(ws, U.faixa(16, P.FORN_INI, 16, P.FORN_FIM))
    nome(wb, "FORNECEDORES_NOMES", f"'{P.ABA_FORN}'!$B${P.FORN_INI}:$B${P.FORN_FIM}")
    nome(wb, "FORNECEDORES_FANTASIA", f"'{P.ABA_FORN}'!$C${P.FORN_INI}:$C${P.FORN_FIM}")
    ws["B3"] = "⚠ DADOS DE EXEMPLO — fornecedores fictícios. Substitua pelos fornecedores reais."
    ws["B3"].style = "NOTA"
    U.preparar(ws, freeze="B6", filtro=f"A5:T{P.FORN_FIM}", repetir_linhas="5:5", tab_cor=E.BLUE)
    ctx.fornecedores = ws
    return ws


# ==========================================================================
# PLANO DE CONTAS
# ==========================================================================
def plano_contas(wb, ctx):
    ws = wb.create_sheet(P.ABA_PLANO)
    headers = ["Código", "Grupo", "Categoria (subgrupo)", "Subcategoria (conta gerencial)",
               "Tipo de conta", "Natureza", "Tipo de fluxo", "Linha do DRE", "Ativo?", "Observações"]
    larg = {"A": 10, "B": 24, "C": 26, "D": 44, "E": 15, "F": 11, "G": 16, "H": 30, "I": 8, "J": 32}
    _cabecalho(ws, "PLANO DE CONTAS GERENCIAL",
               "Estrutura hierárquica que classifica todos os lançamentos e alimenta o DRE, o fluxo de caixa "
               "e o orçamento. Para criar novas contas, use as linhas em branco e escolha os valores das listas.",
               headers, larg, P.PLANO_INI, P.PLANO_FIM)
    for i, p in enumerate(PC.PLANO):
        r = P.PLANO_INI + i
        for j, v in enumerate(p):
            ws.cell(row=r, column=j + 1, value=v)
        ws.cell(row=r, column=9, value="Sim")
        ws.cell(row=r, column=10, value="")
    for r in range(P.PLANO_INI, P.PLANO_FIM + 1):
        for col, est in [(1, "IN_TXT_C"), (2, "IN_TXT"), (3, "IN_TXT"), (4, "IN_TXT"), (5, "IN_TXT_C"),
                         (6, "IN_TXT_C"), (7, "IN_TXT_C"), (8, "IN_TXT"), (9, "IN_TXT_C"), (10, "IN_TXT")]:
            c = ws.cell(row=r, column=col)
            if c.value is None or col > 8:
                c.style = est
            else:
                c.style = est
    rng = f"{P.PLANO_INI}:{P.PLANO_FIM}"
    U.dv_lista(ws, f"E{P.PLANO_INI}:E{P.PLANO_FIM}", "TIPO_CONTA", "Tipo inválido",
               "Escolha um tipo de conta da lista.")
    U.dv_lista(ws, f"F{P.PLANO_INI}:F{P.PLANO_FIM}", "TIPO_LANCAMENTO", "Natureza inválida",
               "Escolha Entrada ou Saída.")
    U.dv_lista(ws, f"G{P.PLANO_INI}:G{P.PLANO_FIM}", "TIPO_FLUXO", "Fluxo inválido",
               "Escolha Operacional, Investimento ou Financiamento.")
    U.dv_lista(ws, f"H{P.PLANO_INI}:H{P.PLANO_FIM}", "LINHAS_DRE", "Linha inválida",
               "Escolha uma linha do DRE da lista.")
    U.dv_lista(ws, f"I{P.PLANO_INI}:I{P.PLANO_FIM}", "SIM_NAO", "Valor inválido", "Escolha Sim ou Não.")
    nome(wb, "CATEGORIAS_PLANO", f"'{P.ABA_PLANO}'!$C${P.PLANO_INI}:$C${P.PLANO_FIM}")
    nome(wb, "SUBGRUPOS_PLANO", f"'{P.ABA_PLANO}'!$C${P.PLANO_INI}:$C${P.PLANO_FIM}")
    nome(wb, "SUBCATEGORIAS_PLANO", f"'{P.ABA_PLANO}'!$D${P.PLANO_INI}:$D${P.PLANO_FIM}")
    nome(wb, "LINHAS_DRE_PLANO", f"'{P.ABA_PLANO}'!$H${P.PLANO_INI}:$H${P.PLANO_FIM}")
    nome(wb, "GRUPOS_PLANO", f"'{P.ABA_PLANO}'!$B${P.PLANO_INI}:$B${P.PLANO_FIM}")
    nome(wb, "TIPOS_FLUXO_PLANO", f"'{P.ABA_PLANO}'!$G${P.PLANO_INI}:$G${P.PLANO_FIM}")
    ws["B3"] = ("Plano de contas de exemplo já cadastrado (receitas, custos, despesas administrativas, pessoal, "
                "comercial, financeiras, tributos, investimentos, financiamentos e patrimônio).")
    ws["B3"].style = "NOTA"
    U.preparar(ws, freeze="A6", filtro=f"A5:J{P.PLANO_FIM}", repetir_linhas="5:5", tab_cor=E.BLUE)
    ctx.plano = ws
    return ws


# ==========================================================================
# CENTRO DE CUSTOS
# ==========================================================================
def centro_custos(wb, ctx):
    ws = wb.create_sheet(P.ABA_CC)
    headers = ["Código", "Centro de custo", "Responsável", "Descrição / finalidade",
               "Ativo?", "Receita (ano)", "Custos (ano)", "Despesas (ano)", "Resultado", "Margem"]
    larg = {"A": 10, "B": 24, "C": 18, "D": 52, "E": 9, "F": 15, "G": 15, "H": 15, "I": 15, "J": 11}
    _cabecalho(ws, "CENTROS DE CUSTO",
               "Separe o resultado da empresa por área. Classifique cada lançamento com um centro de custo "
               "na aba LANÇAMENTOS para ver o desempenho de cada área.", headers, larg, P.CC_INI, P.CC_FIM)
    for i, (cod, nome_cc, resp, desc) in enumerate(D.CENTROS):
        r = P.CC_INI + i
        ws.cell(row=r, column=1, value=cod).style = "IN_TXT_C"
        ws.cell(row=r, column=2, value=nome_cc).style = "IN_TXT"
        ws.cell(row=r, column=3, value=resp).style = "IN_TXT"
        ws.cell(row=r, column=4, value=desc).style = "IN_TXT"
        ws.cell(row=r, column=5, value="Sim").style = "IN_TXT_C"
    for r in range(P.CC_INI, P.CC_FIM + 1):
        for col in range(1, 6):
            if ws.cell(row=r, column=col).value is None:
                ws.cell(row=r, column=col).style = "IN_TXT_C" if col in (1, 5) else "IN_TXT"
        b = f"$B{r}"
        put(ws, r, 6, f'=IF({b}="","",SUMIFS({L_VDRE},{L_CC},{b},{L_TIPO},"Entrada"))', "CALC_MONEY")
        put(ws, r, 7, f'=IF({b}="","",SUMIFS({L_VDRE},{L_CC},{b},{L_LINHA},"Custos Variáveis"))', "CALC_MONEY")
        put(ws, r, 8, f'=IF({b}="","",SUMIFS({L_VDRE},{L_CC},{b},{L_TIPO},"Saída")-'
                      f'SUMIFS({L_VDRE},{L_CC},{b},{L_LINHA},"Custos Variáveis")-'
                      f'SUMIFS({L_VDRE},{L_CC},{b},{L_LINHA},"Investimentos (CAPEX)")-'
                      f'SUMIFS({L_VDRE},{L_CC},{b},{L_LINHA},"Financiamentos e Sócios (Principal)")-'
                      f'SUMIFS({L_VDRE},{L_CC},{b},{L_LINHA},"Distribuição de Lucros"))', "CALC_MONEY")
        put(ws, r, 9, f'=IF({b}="","",$F{r}-$G{r}-$H{r})', "CALC_MONEY")
        put(ws, r, 10, f'=IF(OR({b}="",$F{r}=0),"",$I{r}/$F{r})', "CALC_PCT")
    ult = P.CC_FIM + 1
    put(ws, ult, 2, "TOTAL", "TOTAL_TXT")
    for col in (6, 7, 8, 9):
        put(ws, ult, col, f"=SUM({L(col)}{P.CC_INI}:{L(col)}{P.CC_FIM})", "TOTAL_MONEY")
    put(ws, ult, 10, f'=IF($F{ult}=0,"",$I{ult}/$F{ult})', "TOTAL_PCT")
    U.dv_lista(ws, f"E{P.CC_INI}:E{P.CC_FIM}", "SIM_NAO", "Valor inválido", "Escolha Sim ou Não.")
    nome(wb, "CENTROS_CUSTO", f"'{P.ABA_CC}'!$B${P.CC_INI}:$B${P.CC_FIM}")
    U.cf_escala(ws, f"J{P.CC_INI}:J{P.CC_FIM}")
    U.cf_barras(ws, f"F{P.CC_INI}:F{P.CC_FIM}")
    ws["B3"] = ("⚠ DADOS DE EXEMPLO — 8 centros de custo fictícios (administrativo, comercial, produção, "
                "logística, marketing, financeiro, diretoria e não alocado).")
    ws["B3"].style = "NOTA"

    # ------- gráfico: resultado por centro de custo
    ch = BarChart()
    ch.type = "bar"
    ch.style = 2
    ch.title = "Resultado por centro de custo (ano)"
    data = Reference(ws, min_col=9, min_row=P.CC_INI, max_row=P.CC_INI + len(D.CENTROS) - 1)
    cats = Reference(ws, min_col=2, min_row=P.CC_INI, max_row=P.CC_INI + len(D.CENTROS) - 1)
    ch.add_data(data, titles_from_data=False)
    ch.set_categories(cats)
    ch.y_axis.numFmt = 'R$ #,##0'
    ch.legend = None
    ch.height, ch.width = 7.5, 14
    ws.add_chart(ch, f"L{P.CC_INI}")
    ctx.chart_cc = ch
    U.preparar(ws, freeze="A6", filtro=f"A5:J{P.CC_FIM}", repetir_linhas="5:5", tab_cor=E.BLUE)
    ctx.centro_custos = ws
    return ws


# ==========================================================================
# CONTAS BANCÁRIAS
# ==========================================================================
def contas_bancarias(wb, ctx):
    ws = wb.create_sheet(P.ABA_BANCOS)
    headers = ["Código", "Conta (nome de referência)", "Banco / instituição", "Agência", "Conta nº",
               "Tipo de conta", "Saldo inicial", "Entradas no ano", "Saídas no ano", "Saldo atual",
               "Limite de crédito", "Saldo + limite", "Gerente / contato", "Status", "Observações"]
    larg = {"A": 9, "B": 34, "C": 20, "D": 11, "E": 14, "F": 18, "G": 15, "H": 15, "I": 15, "J": 15,
            "K": 15, "L": 15, "M": 20, "N": 10, "O": 38}
    _cabecalho(ws, "CONTAS BANCÁRIAS",
               "Cadastre as contas correntes e de investimento. O saldo inicial é o saldo existente no "
               "momento em que você começou a usar a planilha; o saldo atual é calculado pelos lançamentos.",
               headers, larg, P.BANCO_INI, P.BANCO_FIM)
    estilos_in = ["IN_TXT_C", "IN_TXT", "IN_TXT", "IN_TXT_C", "IN_TXT_C", "IN_TXT_C", "IN_MONEY",
                  None, None, None, "IN_MONEY", None, "IN_TXT", "IN_TXT_C", "IN_TXT"]
    for i, est in enumerate(estilos_in):
        for r in range(P.BANCO_INI, P.BANCO_FIM + 1):
            if est:
                ws.cell(row=r, column=i + 1).style = est
            else:
                ws.cell(row=r, column=i + 1).style = "CALC_MONEY"
    bancos = [c for c in D.CONTAS if c["tipo"] != "Cartão de crédito"]
    for i, c in enumerate(bancos):
        r = P.BANCO_INI + i
        valores = [f"BN{i+1:02d}", c["nome"], c["banco"], c["agencia"], c["conta"], c["tipo"],
                   c["saldo_inicial"], None, None, None, c["limite"], None, c["gerente"], c["status"], c["obs"]]
        for j, v in enumerate(valores):
            if v is not None:
                ws.cell(row=r, column=j + 1, value=v)
    for r in range(P.BANCO_INI, P.BANCO_FIM + 1):
        b = f"$B{r}"
        put(ws, r, 8, f'=IF({b}="","",SUMIFS({L_VFINAL},{L_BANCO},{b},{L_TIPO},"Entrada",'
                      f'{L_PAGTO},"<="&{HOJE}))', "CALC_MONEY")
        put(ws, r, 9, f'=IF({b}="","",SUMIFS({L_VFINAL},{L_BANCO},{b},{L_TIPO},"Saída",'
                      f'{L_PAGTO},"<="&{HOJE}))', "CALC_MONEY")
        put(ws, r, 10, f'=IF({b}="","",$G{r}+$H{r}-$I{r})', "CALC_MONEY")
        put(ws, r, 12, f'=IF({b}="","",$J{r}+$K{r})', "CALC_MONEY")
    U.dv_lista(ws, f"F{P.BANCO_INI}:F{P.BANCO_FIM}",
               '"Conta Corrente,Conta Poupança,Conta de Investimento,Caixa"',
               "Tipo inválido", "Escolha um tipo de conta.")
    U.dv_lista(ws, f"N{P.BANCO_INI}:N{P.BANCO_FIM}", "STATUS_CADASTRO", "Status inválido",
               "Escolha Ativo, Inativo ou Bloqueado.")
    ult = P.BANCO_FIM + 1
    put(ws, ult, 2, "TOTAL DISPONÍVEL (caixa e bancos)", "TOTAL_TXT")
    for col in (7, 8, 9, 10, 11, 12):
        put(ws, ult, col, f"=SUM({L(col)}{P.BANCO_INI}:{L(col)}{P.BANCO_FIM})", "TOTAL_MONEY")
    nome(wb, "CONTAS_BANCARIAS", f"'{P.ABA_BANCOS}'!$B${P.BANCO_INI}:$B${P.BANCO_FIM}")
    put(ws, ult + 2, 2, "O saldo atual é calculado assim: saldo inicial + entradas realizadas − saídas realizadas "
                        "(data de pagamento/recebimento preenchida na aba LANÇAMENTOS).", "NOTA")
    ws["B3"] = ("⚠ DADOS DE EXEMPLO — três contas correntes fictícias (Banco do Brasil, Itaú e Nubank). "
                "Os cartões de crédito ficam na aba CARTÕES.")
    ws["B3"].style = "NOTA"
    U.preparar(ws, freeze="A6", filtro=f"A5:O{P.BANCO_FIM}", repetir_linhas="5:5", tab_cor=E.BLUE)
    ctx.bancos = ws
    return ws


# ==========================================================================
# CARTÕES
# ==========================================================================
def cartoes(wb, ctx):
    ws = wb.create_sheet(P.ABA_CARTOES)
    headers = ["Código", "Cartão (nome de referência)", "Bandeira", "Banco emissor", "Final",
               "Limite total", "Dia de fechamento", "Dia de vencimento", "Conta de pagamento",
               "Compras no mês (competência)", "Fatura paga no mês", "Fatura em aberto", "Limite disponível",
               "% do limite usado", "Status", "Observações"]
    larg = {"A": 9, "B": 30, "C": 14, "D": 18, "E": 8, "F": 14, "G": 12, "H": 12, "I": 30, "J": 17,
            "K": 16, "L": 14, "M": 15, "N": 12, "O": 10, "P": 40}
    _cabecalho(ws, "CARTÕES DE CRÉDITO",
               "Controle de limites, faturas e comprometimento. As compras devem ser lançadas na aba "
               "LANÇAMENTOS com a forma de pagamento 'Cartão de crédito' e o cartão escolhido na coluna Cartão.",
               headers, larg, P.CARTAO_INI, P.CARTAO_FIM)
    estilos_in = ["IN_TXT_C", "IN_TXT", "IN_TXT_C", "IN_TXT", "IN_TXT_C", "IN_MONEY", "IN_INT", "IN_INT",
                  "IN_TXT", None, None, None, None, None, "IN_TXT_C", "IN_TXT"]
    for i, est in enumerate(estilos_in):
        for r in range(P.CARTAO_INI, P.CARTAO_FIM + 1):
            ws.cell(row=r, column=i + 1).style = est if est else "CALC_MONEY"
    for i, c in enumerate(D.CARTOES):
        r = P.CARTAO_INI + i
        valores = [f"CT{i+1:02d}", c["nome"], c["bandeira"], c["banco"], c["final"], c["limite"],
                   c["fechamento"], c["vencimento"], c["conta_pagamento"], None, None, None, None,
                   None, c["status"], c["obs"]]
        for j, v in enumerate(valores):
            if v is not None:
                ws.cell(row=r, column=j + 1, value=v)
    for r in range(P.CARTAO_INI, P.CARTAO_FIM + 1):
        b = f"$B{r}"
        cfg_ano, cfg_mes = P.CONFIG_REF["ano"], P.CONFIG_REF["mes"]
        mes_comp = P.rng("MES")
        ano_comp = P.rng("ANO")
        put(ws, r, 10, f'=IF({b}="","",SUMIFS({L_VPREV},{L_CARTAO},{b},{mes_comp},{cfg_mes},{ano_comp},{cfg_ano}))',
            "CALC_MONEY")
        put(ws, r, 11, f'=IF({b}="","",SUMIFS({L_VFINAL},{L_CARTAO},{b},{mes_comp},{cfg_mes},{ano_comp},{cfg_ano}))',
            "CALC_MONEY")
        put(ws, r, 12, f'=IF({b}="","",SUMIFS({L_SALDO},{L_CARTAO},{b}))', "CALC_MONEY")
        put(ws, r, 13, f'=IF({b}="","",$F{r}-$L{r})', "CALC_MONEY")
        put(ws, r, 14, f'=IF(OR({b}="",$F{r}=0),"",$L{r}/$F{r})', "CALC_PCT")
    ult = P.CARTAO_FIM + 1
    put(ws, ult, 2, "TOTAL", "TOTAL_TXT")
    for col in (6, 10, 11, 12, 13):
        put(ws, ult, col, f"=SUM({L(col)}{P.CARTAO_INI}:{L(col)}{P.CARTAO_FIM})", "TOTAL_MONEY")
    put(ws, ult, 14, f'=IF($F{ult}=0,"",$L{ult}/$F{ult})', "TOTAL_PCT")
    U.dv_lista(ws, f"C{P.CARTAO_INI}:C{P.CARTAO_FIM}", "BANDEIRA", "Bandeira inválida",
               "Escolha uma bandeira da lista.")
    U.dv_inteiro(ws, f"G{P.CARTAO_INI}:H{P.CARTAO_FIM}", "between", "1", "31", "Dia inválido",
                 "Informe um dia entre 1 e 31.")
    U.dv_lista(ws, f"O{P.CARTAO_INI}:O{P.CARTAO_FIM}", "STATUS_CADASTRO", "Status inválido",
               "Escolha Ativo, Inativo ou Bloqueado.")
    nome(wb, "CARTOES", f"'{P.ABA_CARTOES}'!$B${P.CARTAO_INI}:$B${P.CARTAO_FIM}")
    nome(wb, "CARTAO_LISTA", f"'{P.ABA_CARTOES}'!$B${P.CARTAO_INI}:$B${P.CARTAO_FIM}")

    # ------- faturas por mês (matriz 12 meses)
    linha0 = ult + 3
    U.secao(ws, linha0, "FATURAS POR MÊS — compras em cada cartão (ano de análise)", span=16)
    U.cabecalho(ws, linha0 + 1, 1, ["Cartão"] + [f"Mês {m}" for m in range(1, 13)] + ["Total"], "TH2", 20)
    for i in range(len(D.CARTOES)):
        r = linha0 + 2 + i
        put(ws, r, 1, f"=$B${P.CARTAO_INI + i}", "CALC_TXT")
        for m in range(1, 13):
            put(ws, r, 1 + m,
                f'=IF($A{r}="","",SUMIFS({L_VPREV},{L_CARTAO},$A{r},{P.rng("MES")},{m},{P.rng("ANO")},'
                f'{P.CONFIG_REF["ano"]}))', "CALC_MONEY")
        put(ws, r, 14, f'=IF($A{r}="","",SUM(B{r}:M{r}))', "TOTAL_MONEY")
    U.cf_barras(ws, f"B{linha0+2}:M{linha0+1+len(D.CARTOES)}")
    ws["B3"] = "⚠ DADOS DE EXEMPLO — três cartões de crédito fictícios."
    ws["B3"].style = "NOTA"
    U.preparar(ws, freeze="A6", filtro=f"A5:P{P.CARTAO_FIM}", repetir_linhas="5:5", tab_cor=E.BLUE)
    ctx.cartoes = ws
    return ws
