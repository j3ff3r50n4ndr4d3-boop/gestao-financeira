"""Abas DESPESAS RECORRENTES e PROJEÇÃO DE CAIXA."""
from openpyxl.chart import BarChart, LineChart, Reference

from lib import dados_exemplo as D
from lib import plano_contas as PC

U_GRUPO = {p[3]: p[1] for p in PC.PLANO}
from lib import estilos as E
from lib import parametros as P
from lib import util as U
from lib.util import L, nome, put

L_TIPO = P.rng("TIPO")
L_VENC = P.rng("VENC")
L_COMP = P.rng("COMP")
L_PAGTO = P.rng("PAGTO")
L_SUBCAT = P.rng("SUBCAT")
L_CAT = P.rng("CAT")
L_CC = P.rng("CC")
L_ENT = P.rng("ENTIDADE")
L_DESC = P.rng("DESCRICAO")
L_BANCO = P.rng("BANCO")
L_VPREV = P.rng("VPREV")
L_VREAL = P.rng("VREAL")
L_VFINAL = P.rng("VFINAL")
L_SALDO = P.rng("SALDO")
L_VDRE = P.rng("VDRE")
L_LINHA = P.rng("LINHADRE")
L_MES = P.rng("MES")
L_ANO = P.rng("ANO")
L_FORMA = P.rng("FORMA")

ANO = P.CONFIG_REF["ano"]
MES = P.CONFIG_REF["mes"]
HOJE = "TODAY()"
SALDO_MIN = P.CONFIG_REF["saldo_min"]
DISPONIVEL = f"SUM('{P.ABA_BANCOS}'!$G${P.BANCO_INI}:$G${P.BANCO_FIM})"
SALDO_ATUAL = (f"{DISPONIVEL}+SUMIFS({L_VFINAL},{L_TIPO},\"Entrada\",{L_PAGTO},\"<=\"&{HOJE})"
               f"-SUMIFS({L_VFINAL},{L_TIPO},\"Saída\",{L_PAGTO},\"<=\"&{HOJE})")

MESES_ABREV = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

PERIODOS = {"Mensal": 1, "Bimestral": 2, "Trimestral": 3, "Semestral": 6, "Anual": 12}


def _cards(ws, row, cards, span=2):
    col = 1
    for rotulo, formula, estilo in cards:
        U.cartao(ws, row, col, span, rotulo, formula, estilo)
        col += span
    return row + 3


# ==========================================================================
# DESPESAS RECORRENTES
# ==========================================================================
def despesas_recorrentes(wb, ctx):
    ws = wb.create_sheet(P.ABA_REC)
    U.larguras(ws, {"A": 34, "B": 24, "C": 26, "D": 18, "E": 24, "F": 14, "G": 14, "H": 12, "I": 26,
                    "J": 18, "K": 10, "L": 34, "M": 12, "N": 14, "O": 14, "P": 14, "Q": 14, "R": 14,
                    "S": 14, "T": 14, "U": 14, "V": 14, "W": 14, "X": 14, "Y": 14})
    U.titulo(ws, 1, "DESPESAS RECORRENTES — CADASTRO E GERAÇÃO DE LANÇAMENTOS FUTUROS",
             "Cadastre aluguel, energia, salários, contabilidade e outras despesas fixas. O sistema calcula as "
             "próximas ocorrências, gera a lista pronta para lançamento e projeta os próximos 12 meses.", span=8)
    U.mesclar(ws, "A3:L3",
              "Como usar: 1) cadastre a despesa (células amarelas); 2) use a lista do bloco 3 para lançar as "
              "ocorrências na aba LANÇAMENTOS; 3) as ocorrências já lançadas são desconsideradas "
              "automaticamente no bloco 4 (evitando duplicidade). ⚠ As despesas cadastradas são exemplos.",
              "NOTA_WRAP")
    ws.row_dimensions[3].height = 30

    ini, fim = P.RECOR_INI, P.RECOR_FIM
    U.secao(ws, 5, "1. CADASTRO DAS DESPESAS RECORRENTES", span=8)
    U.cabecalho(ws, 6, 1, ["Descrição da despesa", "Categoria", "Subcategoria", "Centro de custo",
                           "Fornecedor / prestador", "Valor mensal", "Periodicidade", "Dia do venc.",
                           "Banco / conta", "Forma de pagamento", "Ativa?", "Observações"], "TH", 34)
    for r in range(ini, fim + 1):
        for col, est in [(1, "IN_TXT"), (2, "IN_TXT_C"), (3, "IN_TXT_C"), (4, "IN_TXT_C"), (5, "IN_TXT"),
                         (6, "IN_MONEY"), (7, "IN_TXT_C"), (8, "IN_INT"), (9, "IN_TXT_C"), (10, "IN_TXT_C"),
                         (11, "IN_TXT_C"), (12, "IN_TXT")]:
            ws.cell(row=r, column=col).style = est
    for i, rc in enumerate(D.RECORRENTES):
        r = ini + i
        for col, v in [(1, rc["desc"]), (2, rc.get("cat", U_GRUPO.get(rc["subcat"], "Despesas Operacionais"))), (3, rc["subcat"]), (4, rc["cc"]), (5, rc.get("entidade", "")),
                       (6, rc["valor"]), (7, rc.get("periodicidade", "Mensal")), (8, rc["dia"]),
                       (9, rc["banco"]), (10, rc["forma"]), (11, rc.get("ativa", "Sim")),
                       (12, rc.get("obs", ""))]:
            if v not in (None, ""):
                ws.cell(row=r, column=col, value=v)
    U.dv_lista(ws, f"B{ini}:B{fim}", "SUBGRUPOS_PLANO", "Categoria inválida",
               "Escolha uma categoria do plano de contas.")
    U.dv_lista(ws, f"C{ini}:C{fim}", "SUBCATEGORIAS_PLANO", "Subcategoria inválida",
               "Escolha uma subcategoria do plano de contas.", estilo="warning")
    U.dv_lista(ws, f"D{ini}:D{fim}", "CENTROS_CUSTO", "Centro de custo inválido",
               "Escolha um centro de custo cadastrado.", estilo="warning")
    U.dv_lista(ws, f"G{ini}:G{fim}", "PERIODICIDADE", "Periodicidade inválida",
               "Escolha Mensal, Bimestral, Trimestral, Semestral ou Anual.")
    U.dv_lista(ws, f"I{ini}:I{fim}", "CONTAS_BANCARIAS", "Conta inválida",
               "Escolha uma conta bancária ou cartão cadastrado.", estilo="warning")
    U.dv_lista(ws, f"J{ini}:J{fim}", "FORMA_PAGAMENTO", "Forma inválida",
               "Escolha uma forma de pagamento da lista.")
    U.dv_lista(ws, f"K{ini}:K{fim}", "SIM_NAO", "Valor inválido", "Use Sim ou Não.")
    U.dv_inteiro(ws, f"H{ini}:H{fim}", "between", "1", "31", "Dia inválido",
                 "Informe o dia do vencimento (1 a 31).")
    U.dv_decimal(ws, f"F{ini}:F{fim}", "greaterThanOrEqual", "0", "Valor inválido",
                 "Informe o valor mensal da despesa.")

    # indicadores
    r_ind = fim + 2
    U.secao(ws, r_ind, "2. INDICADORES DAS DESPESAS RECORRENTES", span=8)
    tot_mensal = f'SUMIFS($F${ini}:$F${fim},$K${ini}:$K${fim},"Sim")'
    cards = [
        ("Total mensal das recorrentes", f"={tot_mensal}", "KPI_VALUE"),
        ("Total anual projetado", f"={tot_mensal}*12", "KPI_VALUE"),
        ("% das despesas fixas do mês",
         f'=IFERROR({tot_mensal}/MAX(1,SUMIFS({L_VDRE},{L_LINHA},"Despesas Fixas",{L_ANO},{ANO},'
         f'{L_MES},{MES})),"")', "KPI_PCT"),
        ("Despesas recorrentes cadastradas", f'=COUNTIFS($K${ini}:$K${fim},"Sim")', "KPI_VALUE0"),
    ]
    _cards(ws, r_ind + 1, cards, span=2)
    U.cf_formula(ws, f"F{r_ind+2}", f'=F{r_ind+2}>0', None, E.RED_TX)

    # ------------------------------------------- lista pronta para lançar
    r_lst = r_ind + 6
    U.secao(ws, r_lst, "3. PRÓXIMAS OCORRÊNCIAS — LISTA PRONTA PARA LANÇAR NA ABA LANÇAMENTOS", span=8)
    U.cabecalho(ws, r_lst + 1, 1, ["Descrição", "Data de competência", "Data de vencimento", "Tipo",
                                   "Categoria", "Subcategoria", "Centro de custo", "Fornecedor",
                                   "Valor previsto", "Banco / conta", "Forma de pagamento",
                                   "Já lançado no mês?"], "TH2", 34)
    r0 = r_lst + 2
    for i in range(fim - ini + 1):
        r = r0 + i
        ref = ini + i
        base = f"${ref}"
        # vencimento: se o dia já passou neste mês, agenda para o próximo mês
        venc = (f'IF(DATE(YEAR({HOJE}),MONTH({HOJE}),MIN($H{base},DAY(EOMONTH({HOJE},0))))<{HOJE},'
                f'DATE(YEAR(EDATE({HOJE},1)),MONTH(EDATE({HOJE},1)),MIN($H{base},'
                f'DAY(EOMONTH(EDATE({HOJE},1),0)))),'
                f'DATE(YEAR({HOJE}),MONTH({HOJE}),MIN($H{base},DAY(EOMONTH({HOJE},0)))))')
        put(ws, r, 1, f'=IF(OR($A{base}="",$K{base}<>"Sim"),"",$A{base})', "CALC_TXT")
        put(ws, r, 2, f'=IF($A{r}="","",EOMONTH($C{r},-1)+1)', "CALC_DATE")
        put(ws, r, 3, f'=IF($A{r}="","",{venc})', "CALC_DATE")
        put(ws, r, 4, "Saída", "CALC_TXT_C")
        put(ws, r, 5, f'=IF($A{r}="","",$B{base})', "CALC_TXT_C")
        put(ws, r, 6, f'=IF($A{r}="","",$C{base})', "CALC_TXT_C")
        put(ws, r, 7, f'=IF($A{r}="","",$D{base})', "CALC_TXT_C")
        put(ws, r, 8, f'=IF($A{r}="","",$E{base})', "CALC_TXT")
        put(ws, r, 9, f'=IF($A{r}="","",$F{base})', "CALC_MONEY")
        put(ws, r, 10, f'=IF($A{r}="","",$I{base})', "CALC_TXT_C")
        put(ws, r, 11, f'=IF($A{r}="","",$J{base})', "CALC_TXT_C")
        put(ws, r, 12, f'=IF($A{r}="","",IF(COUNTIFS({L_SUBCAT},$F{r},{L_TIPO},"Saída",{L_MES},'
                       f'MONTH($B{r}),{L_ANO},YEAR($B{r}))>0,"✔ Sim","— Não"))', "CALC_TXT_C")
    rt = r0 + (fim - ini)
    put(ws, rt, 1, "TOTAL A LANÇAR", "TOTAL_TXT")
    put(ws, rt, 9, f'=SUMIF($L${r0}:$L${rt-1},"— Não",$I${r0}:$I${rt-1})', "TOTAL_MONEY")
    U.cf_contem(ws, f"L{r0}:L{rt-1}", "✔", E.GREEN_BG, E.GREEN_TX)
    U.cf_contem(ws, f"L{r0}:L{rt-1}", "Não", E.YEL_BG, E.YEL_TX)

    # ------------------------------------------- projeção 12 meses
    r_prj = rt + 2
    U.secao(ws, r_prj, "4. PROJEÇÃO DAS RECORRENTES PARA OS PRÓXIMOS 12 MESES (desconsidera o que já "
                       "está lançado)", span=8, altura=22)
    cab = ["Descrição", "Subcategoria"]
    meses = []
    for k in range(12):
        d = f"EDATE({HOJE},{k})"
        meses.append(d)
        cab.append(f'={d}')
    U.cabecalho(ws, r_prj + 1, 1, ["Descrição", "Subcategoria"] + [f"Mês {k+1}" for k in range(12)],
                "TH2", 30)
    # os rótulos abaixo são substituídos por DATAS (1º dia de cada mês): a PROJEÇÃO DE CAIXA
    # compara esta linha com as datas de vencimento/recorrência
    for k in range(12):
        # 1º dia do mês de referência: a PROJEÇÃO DE CAIXA compara esta linha com o horizonte
        d = f"EDATE({HOJE},{k})"
        c = ws.cell(row=r_prj + 1, column=3 + k,
                    value=f"=DATE(YEAR({d}),MONTH({d}),1)")
        c.style = "TH2"
        c.number_format = "MM/AAAA"
    ws.row_dimensions[r_prj + 1].height = 30
    r0p = r_prj + 2
    for i in range(fim - ini + 1):
        r = r0p + i
        ref = ini + i
        base = f"${ref}"
        put(ws, r, 1, f'=IF(OR($A{base}="",$K{base}<>"Sim"),"",$A{base})', "CALC_TXT")
        put(ws, r, 2, f'=IF($A{r}="","",$C{base})', "CALC_TXT_C")
        for k in range(12):
            col = 3 + k
            d = f"EDATE({HOJE},{k})"
            n = (f'IF($G{base}="Mensal",1,IF($G{base}="Bimestral",2,IF($G{base}="Trimestral",3,'
                 f'IF($G{base}="Semestral",6,12))))')
            f = (f'=IF($A{r}="",0,IF(MOD({k},{n})<>0,0,'
                 f'IF(COUNTIFS({L_SUBCAT},$B{r},{L_TIPO},"Saída",{L_MES},MONTH({d}),{L_ANO},YEAR({d}))>0,0,'
                 f'$F{base})))')
            put(ws, r, col, f, "CALC_MONEY")
    rtp = r0p + (fim - ini)
    put(ws, rtp, 1, "TOTAL DAS RECORRENTES NO MÊS", "TOTAL_TXT")
    put(ws, rtp, 2, "", "TOTAL_TXT")
    for k in range(12):
        col = 3 + k
        put(ws, rtp, col, f"=SUM({L(col)}{r0p}:{L(col)}{rtp-1})", "TOTAL_MONEY")
    U.cf_barras(ws, f"C{rtp}:N{rtp}")
    nome(wb, "RECORRENTES_MES", f"'{P.ABA_REC}'!$C${rtp}:$N${rtp}")
    ctx.rec_grid = dict(r0=r0p, r1=rtp - 1, total=rtp, meses=r_prj + 1)

    ch = BarChart()
    ch.type = "col"
    ch.title = "Despesas recorrentes projetadas (próximos 12 meses)"
    ch.style = 2
    ch.add_data(Reference(ws, min_col=3, max_col=14, min_row=rtp, max_row=rtp), from_rows=True,
                titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=3, max_col=14, min_row=r_prj + 1, max_row=r_prj + 1))
    ch.y_axis.numFmt = 'R$ #,##0'
    ch.legend = None
    ch.height, ch.width = 8, 24
    ws.add_chart(ch, f"A{rtp + 2}")
    ctx.chart_recorrentes = ch
    ctx.recorrentes = ws
    U.preparar(ws, freeze="A7", filtro=f"A6:L{fim}", tab_cor="808080",
               area_impressao=f"A1:L{rt+2}")
    return ws


# ==========================================================================
# PROJEÇÃO DE CAIXA
# ==========================================================================
def projecao_caixa(wb, ctx):
    ws = wb.create_sheet(P.ABA_PROJ)
    U.larguras(ws, {"A": 42, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 20, "H": 14, "I": 14,
                    "J": 14, "K": 14, "L": 14, "M": 14})
    U.titulo(ws, 1, "PROJEÇÃO DE CAIXA — 7, 15, 30, 60, 90 E 180 DIAS",
             "Simulação do saldo de caixa no curto e médio prazo, considerando os títulos em aberto, as "
             "despesas recorrentes cadastradas e as parcelas de empréstimos.", span=8)
    U.mesclar(ws, "A3:H3",
              "Como ler: as entradas previstas usam os títulos a receber em aberto pela data de vencimento; "
              "as saídas usam os títulos a pagar em aberto + as despesas recorrentes ainda não lançadas. "
              "O alerta de déficit compara o saldo projetado com o saldo mínimo definido nas configurações.",
              "NOTA_WRAP")
    ws.row_dimensions[3].height = 30

    U.secao(ws, 5, "1. POSIÇÃO ATUAL", span=8)
    cards = [
        ("Saldo disponível hoje", f"={SALDO_ATUAL}", "KPI_VALUE"),
        ("A receber em aberto", f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada")', "KPI_VALUE"),
        ("A pagar em aberto", f'=SUMIFS({L_SALDO},{L_TIPO},"Saída")', "KPI_VALUE"),
        ("Saldo mínimo exigido", f"={SALDO_MIN}", "KPI_VALUE"),
    ]
    _cards(ws, 6, cards, span=2)
    resultado = f'=B$7+SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_VENC},{L_VENC},">="&{HOJE},{{venc}})-SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_VENC},{L_VENC},">="&{HOJE},{{venc}})'
    U.secao(ws, 9, "2. PROJEÇÃO POR HORIZONTE DE TEMPO", span=8)
    U.cabecalho(ws, 10, 1, ["Horizonte", "Entradas previstas", "Saídas previstas", "Recorrentes não lançadas",
                            "Resultado do período", "Saldo projetado", "Saldo mínimo exigido",
                            "Folga / déficit", "Situação"], "TH2", 34)
    r0 = 11
    horizontes = [7, 15, 30, 60, 90, 180]
    for i, h in enumerate(horizontes):
        r = r0 + i
        put(ws, r, 1, f'="{h} dias"', "CALC_TXT_C")
        put(ws, r, 2, f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+{h})',
            "CALC_MONEY")
        put(ws, r, 3, f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+{h})',
            "CALC_MONEY")
        put(ws, r, 4,
            f'=SUMPRODUCT((\'{P.ABA_REC}\'!$C${ctx.rec_grid["meses"]}:$N${ctx.rec_grid["meses"]}'
            f'<={HOJE}+{h})*(\'{P.ABA_REC}\'!$C${ctx.rec_grid["meses"]}:$N${ctx.rec_grid["meses"]}'
            f'>EOMONTH({HOJE},0))*\'{P.ABA_REC}\'!$C${ctx.rec_grid["total"]}:$N${ctx.rec_grid["total"]})',
            "CALC_MONEY")
        put(ws, r, 5, f'=$B{r}-$C{r}-$D{r}', "CALC_MONEY")
        put(ws, r, 6, f'=$A$7+$E{r}', "CALC_MONEY")
        put(ws, r, 7, f"={SALDO_MIN}", "CALC_MONEY")
        put(ws, r, 8, f'=$F{r}-$G{r}', "CALC_MONEY")
        put(ws, r, 9, f'=IF($H{r}<0,"🔴 Déficit de caixa projetado","🟢 Caixa suficiente no período")',
            "CALC_TXT_C")
    U.cf_contem(ws, f"I{r0}:I{r0+5}", "🟢", E.GREEN_BG, E.GREEN_TX)
    U.cf_contem(ws, f"I{r0}:I{r0+5}", "🔴", E.RED_BG, E.RED_TX)
    U.cf_maior(ws, f"H{r0}:H{r0+5}", 0, None, E.RED_TX, op="lessThan")
    U.cf_barras(ws, f"F{r0}:F{r0+5}")

    # -------------------------------------------------- mensal 12 meses
    r_m = r0 + 7
    U.secao(ws, r_m, "3. PROJEÇÃO MENSAL — PRÓXIMOS 12 MESES", span=8)
    U.cabecalho(ws, r_m + 1, 1, ["Mês", "Entradas previstas", "Saídas previstas (títulos)",
                                 "Despesas recorrentes não lançadas", "Resultado do mês",
                                 "Saldo projetado acumulado", "Saldo mínimo", "Folga / déficit",
                                 "Situação", "Receita estimada usada"], "TH2", 34)
    r0m = r_m + 2
    for k in range(12):
        r = r0m + k
        dm = f"EDATE({HOJE},{k})"
        ini_m = f"DATE(YEAR({dm}),MONTH({dm}),1)"
        fim_m = f"EOMONTH({dm},0)"
        put(ws, r, 1, f'=TEXT({dm},"MM/AAAA")', "CALC_TXT_C")
        put(ws, r, 2, f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_VENC},">="&{ini_m},{L_VENC},"<="&{fim_m})',
            "CALC_MONEY")
        put(ws, r, 3, f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_VENC},">="&{ini_m},{L_VENC},"<="&{fim_m})',
            "CALC_MONEY")
        put(ws, r, 4,
            f'=SUMPRODUCT((\'{P.ABA_REC}\'!$C${ctx.rec_grid["meses"]}:$N${ctx.rec_grid["meses"]}'
            f'>={ini_m})*(\'{P.ABA_REC}\'!$C${ctx.rec_grid["meses"]}:$N${ctx.rec_grid["meses"]}'
            f'<={fim_m})*\'{P.ABA_REC}\'!$C${ctx.rec_grid["total"]}:$N${ctx.rec_grid["total"]})',
            "CALC_MONEY")
        put(ws, r, 5, f'=-$C{r}-$D{r}+IF($B{r}=0,$J{r},$B{r})', "CALC_MONEY")
        if k == 0:
            put(ws, r, 6, f'=$A$7+$E{r}', "CALC_MONEY")
        else:
            put(ws, r, 6, f'=$F{r-1}+$E{r}', "CALC_MONEY")
        put(ws, r, 7, f"={SALDO_MIN}", "CALC_MONEY")
        put(ws, r, 8, f'=$F{r}-$G{r}', "CALC_MONEY")
        put(ws, r, 9, f'=IF($H{r}<0,"🔴 Déficit em "&$A{r},"🟢 Caixa suficiente")', "CALC_TXT_C")
        put(ws, r, 10, f'=IF($B{r}>0,0,ROUND(AVERAGEIF(\'{P.ABA_IND}\'!$B${ctx.ind_trend[0]}:$B$'
                       f'{ctx.ind_trend[1]},">0"),2))', "CALC_MONEY")
    r_fim = r0m + 11
    put(ws, r_fim + 1, 1, "TOTAL DOS 12 MESES", "TOTAL_TXT")
    for col in (2, 3, 4, 5):
        put(ws, r_fim + 1, col, f"=SUM({L(col)}{r0m}:{L(col)}{r_fim})", "TOTAL_MONEY")
    put(ws, r_fim + 1, 6, f"=$F{r_fim}", "TOTAL_MONEY")
    put(ws, r_fim + 1, 8, f"=$H{r_fim}", "TOTAL_MONEY")
    U.cf_contem(ws, f"I{r0m}:I{r_fim}", "🟢", E.GREEN_BG, E.GREEN_TX)
    U.cf_contem(ws, f"I{r0m}:I{r_fim}", "🔴", E.RED_BG, E.RED_TX)
    U.cf_barras(ws, f"F{r0m}:F{r_fim}")

    ch = LineChart()
    ch.title = "Saldo de caixa projetado (12 meses)"
    ch.style = 2
    ch.add_data(Reference(ws, min_col=6, min_row=r_m + 1, max_row=r_fim), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0m, max_row=r_fim))
    ch.y_axis.numFmt = 'R$ #,##0'
    ch.height, ch.width = 8, 16
    ws.add_chart(ch, f"J{r0}")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.grouping = "clustered"
    ch2.title = "Entradas x saídas previstas (12 meses)"
    ch2.style = 2
    ch2.add_data(Reference(ws, min_col=2, max_col=3, min_row=r_m + 1, max_row=r_fim), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=r0m, max_row=r_fim))
    ch2.y_axis.numFmt = 'R$ #,##0'
    ch2.height, ch2.width = 8, 16
    ws.add_chart(ch2, f"J{r0+17}")

    ctx.chart_projecao = ch
    ctx.projecao = ws
    U.preparar(ws, tab_cor=E.BLUE, area_impressao=f"A1:I{r_fim+2}")
    return ws
