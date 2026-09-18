"""Abas PAINEL DE CONTROLE, DASHBOARD, RELATÓRIOS e INSTRUÇÕES."""
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Alignment, Font

from lib import dados_exemplo as D
from lib import estilos as E
from lib import parametros as P
from lib import plano_contas as PC
from lib import util as U
from lib.util import L, nome, put

L_TIPO = P.rng("TIPO")
L_VENC = P.rng("VENC")
L_PAGTO = P.rng("PAGTO")
L_SUBCAT = P.rng("SUBCAT")
L_CAT = P.rng("CAT")
L_CC = P.rng("CC")
L_BANCO = P.rng("BANCO")
L_VPREV = P.rng("VPREV")
L_VFINAL = P.rng("VFINAL")
L_SALDO = P.rng("SALDO")
L_VDRE = P.rng("VDRE")
L_LINHA = P.rng("LINHADRE")
L_MES = P.rng("MES")
L_ANO = P.rng("ANO")

ANO = P.CONFIG_REF["ano"]
MES = P.CONFIG_REF["mes"]
HOJE = "TODAY()"
RAZAO = P.CONFIG_REF["razao"]
FANTASIA = P.CONFIG_REF["fantasia"]
SALDO_MIN = P.CONFIG_REF["saldo_min"]
META_FAT = P.CONFIG_REF["meta_fat"]
MARGEM_MIN = P.CONFIG_REF["margem_min"]
DISPONIVEL = f"SUM('{P.ABA_BANCOS}'!$G${P.BANCO_INI}:$G${P.BANCO_FIM})"

LINHAS_DRE = ["Faturamento Bruto", "Impostos sobre Vendas", "Devoluções", "Custos Variáveis",
              "Despesas Fixas", "Despesas Financeiras", "Receitas Financeiras",
              "Outras Receitas Operacionais", "Impostos sobre o Resultado", "Distribuição de Lucros",
              "Investimentos (CAPEX)", "Financiamentos e Sócios (Principal)"]


# ------------------------------------------------------------------ gráficos
def _série(cor):
    from openpyxl.chart.shapes import GraphicalProperties
    return GraphicalProperties(solidFill=cor)


def _base(ch, titulo, altura=7.4, largura=12.6):
    ch.title = titulo
    ch.style = 2
    ch.height, ch.width = altura, largura
    ch.y_axis.numFmt = 'R$ #,##0'
    ch.y_axis.majorGridlines = None
    return ch


def _grafico_col(ws_src, col_dad, col_cat, r_ini, r_fim, titulo, cor=E.NAVY, altura=7.4,
                 largura=12.6, fmt='R$ #,##0', header=True):
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    d = Reference(ws_src, min_col=col_dad, min_row=r_ini - 1 if header else r_ini, max_row=r_fim)
    ch.add_data(d, titles_from_data=header)
    ch.set_categories(Reference(ws_src, min_col=col_cat, min_row=r_ini, max_row=r_fim))
    _base(ch, titulo, altura, largura)
    ch.y_axis.numFmt = fmt
    ch.legend = None
    if ch.series:
        ch.series[0].graphicalProperties = _série(cor)
    return ch


def _grafico_linha(ws_src, col_dad, col_cat, r_ini, r_fim, titulo, cor=E.BLUE, altura=7.4,
                   largura=12.6, fmt='R$ #,##0', header=True):
    ch = LineChart()
    d = Reference(ws_src, min_col=col_dad, min_row=r_ini - 1 if header else r_ini, max_row=r_fim)
    ch.add_data(d, titles_from_data=header)
    ch.set_categories(Reference(ws_src, min_col=col_cat, min_row=r_ini, max_row=r_fim))
    _base(ch, titulo, altura, largura)
    ch.y_axis.numFmt = fmt
    ch.legend = None
    if ch.series:
        ch.series[0].graphicalProperties = _série(cor)
        ch.series[0].smooth = False
    return ch


def _grafico_multiplo(ws_src, cols, col_cat, r_ini, r_fim, titulo, cores=(E.NAVY, E.BLUE),
                      altura=7.4, largura=12.6, fmt='R$ #,##0', tipo="col"):
    ch = BarChart() if tipo == "col" else LineChart()
    if tipo == "col":
        ch.type = "col"
        ch.grouping = "clustered"
    d = Reference(ws_src, min_col=min(cols), max_col=max(cols), min_row=r_ini - 1, max_row=r_fim)
    ch.add_data(d, titles_from_data=True)
    ch.set_categories(Reference(ws_src, min_col=col_cat, min_row=r_ini, max_row=r_fim))
    _base(ch, titulo, altura, largura)
    ch.y_axis.numFmt = fmt
    for i, s in enumerate(ch.series):
        s.graphicalProperties = _série(cores[i % len(cores)])
    return ch


def _grafico_pizza(ws_src, col_dad, col_cat, r_ini, r_fim, titulo, doughnut=False, altura=7.4,
                   largura=12.6):
    ch = DoughnutChart() if doughnut else PieChart()
    d = Reference(ws_src, min_col=col_dad, min_row=r_ini - 1, max_row=r_fim)
    ch.add_data(d, titles_from_data=True)
    ch.set_categories(Reference(ws_src, min_col=col_cat, min_row=r_ini, max_row=r_fim))
    ch.title = titulo
    ch.height, ch.width = altura, largura
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showPercent = True
    return ch


def _cartao_filtro(ws, row, col, span, label, valor, formato="KPI_VALUE"):
    U.cartao(ws, row, col, span, label, valor, formato)


# ==========================================================================
# DASHBOARD
# ==========================================================================
def dashboard(wb, ctx):
    ws = wb.create_sheet(P.ABA_DASH)
    ws.sheet_view.showGridLines = False
    U.larguras(ws, {"A": 2.5, "R": 2.5})
    for col in "BCDEFGHIJKLMNOPQ":
        ws.column_dimensions[col].width = 11.6

    # cabeçalho
    U.mesclar(ws, "B1:Q2", f'=IF({FANTASIA}="","DASHBOARD EXECUTIVO",{FANTASIA}&" — DASHBOARD EXECUTIVO")',
              "TITULO_DASH")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 14
    U.mesclar(ws, "B3:Q3",
              f'="Atualizado em "&TEXT({HOJE},"DD/MM/YYYY")&"   |   Ano de análise: "&{ANO}&"   |   '
              f'Mês de análise: "&IF({MES}=0,"ano completo",TEXT(DATE({ANO},{MES},1),"MM/AAAA"))&'
              f'"   |   Todos os valores são calculados automaticamente a partir da aba LANÇAMENTOS."',
              "SUBTITULO_DASH")

    # ---------------------------------------------------------------- filtros
    U.secao(ws, 5, "🎛️ FILTROS — altere para analisar o período, o centro de custo, a categoria, "
                   "a conta ou o tipo de operação", span=16, col=2, altura=22)
    filtros = [
        ("Ano", f"={ANO}", "IN_INT", "anos"),
        ("Mês (0 = ano todo)", f"={MES}", "IN_INT", "meses"),
        ("Centro de custo", "Todos", "IN_TXT_C", "centros"),
        ("Categoria (subcategoria)", "Todos", "IN_TXT_C", "categorias"),
        ("Conta / cartão", "Todos", "IN_TXT_C", "contas"),
        ("Tipo de operação", "Todos", "IN_TXT_C", "tipos"),
    ]
    for i, (rot, val, est, tipo) in enumerate(filtros):
        col = 2 + i * 2
        U.mesclar(ws, f"{L(col)}6:{L(col+1)}6", rot, "FILTRO_LABEL")
        U.mesclar(ws, f"{L(col)}7:{L(col+1)}7", val, est)
    U.dv_inteiro(ws, "B7:C7", "between", "2000", "2100", "Ano inválido", "Informe um ano de 4 dígitos.")
    U.dv_inteiro(ws, "D7:E7", "between", "0", "12", "Mês inválido", "Use 0 para o ano completo ou 1 a 12.")
    U.dv_lista(ws, "F7:G7", "CENTROS_CUSTO", "Centro de custo inválido",
               "Escolha um centro de custo cadastrado ou digite Todos.", estilo="warning")
    U.dv_lista(ws, "H7:I7", "SUBCATEGORIAS_PLANO", "Categoria inválida",
               "Escolha uma categoria do plano de contas ou digite Todos.", estilo="warning")
    U.dv_lista(ws, "J7:K7", "CONTAS_BANCARIAS", "Conta inválida",
               "Escolha uma conta/cartão cadastrado ou digite Todos.", estilo="warning")
    U.dv_lista(ws, "L7:M7", "TIPO_LANCAMENTO", "Tipo inválido",
               "Escolha Entrada, Saída ou digite Todos.", estilo="warning")
    U.mesclar(ws, "N6:Q7", '="Lançamentos no filtro: "&TEXT($B$146,"#,##0")', "KPI_VALUE_TXT")

    # ------------------------------------------------------------- indicadores
    U.secao(ws, 10, "📌 INDICADORES PRINCIPAIS DO PERÍODO FILTRADO", span=16, col=2, altura=22)
    cards1 = [
        ("Faturamento do período", "=$B$131", "KPI_VALUE"),
        ("Lucro líquido do período", "=$B$140", "KPI_VALUE"),
        ("Margem líquida", "=$B$141", "KPI_PCT"),
        ("Saldo de caixa hoje", f'={DISPONIVEL}+SUMIFS({L_VFINAL},{L_TIPO},"Entrada",{L_PAGTO},"<="&{HOJE})'
                                f'-SUMIFS({L_VFINAL},{L_TIPO},"Saída",{L_PAGTO},"<="&{HOJE})', "KPI_VALUE"),
        ("A receber em aberto", "=$B$147", "KPI_VALUE"),
        ("A pagar em aberto", "=$B$148", "KPI_VALUE"),
        ("Vencido a receber", "=$B$149", "KPI_VALUE"),
        ("Vencido a pagar", "=$B$150", "KPI_VALUE"),
    ]
    for i, (rot, val, est) in enumerate(cards1):
        col = 2 + i * 2
        _cartao_filtro(ws, 11, col, 2, rot, val, est)
    cards2 = [
        ("Entradas de caixa do período", "=$B$142", "KPI_VALUE"),
        ("Saídas de caixa do período", "=$B$143", "KPI_VALUE"),
        ("Geração de caixa do período", "=$B$144", "KPI_VALUE"),
        ("Orçado no período", "=$B$151", "KPI_VALUE"),
        ("Realizado no período", "=$B$152", "KPI_VALUE"),
        ("Consumo do orçamento", "=$B$153", "KPI_PCT"),
        ("Margem de contribuição", "=$B$139", "KPI_PCT"),
        ("Lançamentos no filtro", "=$B$146", "KPI_VALUE0"),
    ]
    for i, (rot, val, est) in enumerate(cards2):
        col = 2 + i * 2
        _cartao_filtro(ws, 13, col, 2, rot, val, est)
    U.cf_formula(ws, "F12:G12", f'=AND($B$141<>"",$B$141<{MARGEM_MIN})', E.RED_BG, E.RED_TX, True)
    U.cf_formula(ws, "F14:G14", '=$B$144<0', E.RED_BG, E.RED_TX, True)
    U.cf_formula(ws, "H12:I12", f'=$H$12<{P.CONFIG_REF["saldo_min"]}', E.RED_BG, E.RED_TX, True)
    U.cf_formula(ws, "N12:O12", '=$N$12>0', E.RED_BG, E.RED_TX, True)
    U.cf_formula(ws, "P12:Q12", '=$P$12>0', E.YEL_BG, E.YEL_TX, True)

    # -------------------------------------------------------------- alertas
    U.secao(ws, 16, "🔔 ATENÇÃO DO GESTOR — SITUAÇÕES QUE EXIGEM DECISÃO", span=16, col=2, altura=22)
    r_al = ctx.ind_alertas
    for i in range(9):
        # esquerda (alertas 1-9)
        r = 17 + i
        U.mesclar(ws, f"B{r}:H{r}", f"='{P.ABA_IND}'!$A${r_al[0] + i}", "ALERTA_TXT")
        U.mesclar(ws, f"I{r}:K{r}", f"='{P.ABA_IND}'!$B${r_al[0] + i}", "CALC_TXT_C")
        U.mesclar(ws, f"L{r}:M{r}", f"='{P.ABA_IND}'!$C${r_al[0] + i}", "CALC_MONEY")
        # direita (alertas 10-18)
        if r_al[0] + 9 + i <= r_al[1]:
            U.mesclar(ws, f"N{r}:P{r}", f"='{P.ABA_IND}'!$A${r_al[0] + 9 + i}", "ALERTA_TXT")
            U.mesclar(ws, f"Q{r}:Q{r}", f"='{P.ABA_IND}'!$B${r_al[0] + 9 + i}", "ALERTA_ICONE")
    ws.row_dimensions[16].height = 22
    for i in range(9):
        ws.row_dimensions[17 + i].height = 18
    rng_al = "I17:K25"
    U.cf_contem(ws, rng_al, "🔴", E.RED_BG, E.RED_TX)
    U.cf_contem(ws, rng_al, "🟡", E.YEL_BG, E.YEL_TX)
    U.cf_contem(ws, rng_al, "🟢", E.GREEN_BG, E.GREEN_TX)
    U.mesclar(ws, "B26:Q26", "Painel completo com as 18 verificações, valores de referência e ações "
                             "recomendadas na aba INDICADORES (bloco ATENÇÃO DO GESTOR).", "NOTA_WRAP")

    # ------------------------------------------------------- resumo executivo
    U.secao(ws, 27, "📋 RESUMO EXECUTIVO DO PERÍODO FILTRADO", span=16, col=2, altura=22)
    resumo = [
        ("Faturamento do período", "=$B$131", "KPI_VALUE"), ("Impostos e devoluções", "=$B$132", "KPI_VALUE"),
        ("Receita líquida", "=$B$136", "KPI_VALUE"), ("Custos variáveis", "=$B$133", "KPI_VALUE"),
        ("Despesas fixas", "=$B$134", "KPI_VALUE"), ("Despesas financeiras", "=$B$135", "KPI_VALUE"),
        ("Impostos sobre o lucro", "=$B$138", "KPI_VALUE"), ("Lucro líquido", "=$B$140", "KPI_VALUE"),
    ]
    for i, (rot, val, est) in enumerate(resumo):
        r = 28 + (i // 4)
        col = 2 + (i % 4) * 4
        U.mesclar(ws, f"{L(col)}{r}:{L(col+1)}{r}", rot, "KPI_LABEL")
        U.mesclar(ws, f"{L(col+2)}{r}:{L(col+3)}{r}", val, est)

    # ------------------------------------------------------------- gráficos
    U.secao(ws, 31, "📈 GRÁFICOS GERENCIAIS — atualizam automaticamente com os lançamentos", span=16,
            col=2, altura=22)
    ind = ctx.indicadores
    ti, tf = ctx.ind_trend
    dc = ctx.dre_cat
    dre = ctx.dre
    apf = ctx.ap_faixa
    arf = ctx.ar_faixa
    inad = ctx.inad_faixa
    er = ctx.emp_rows
    emp = ctx.emprestimos
    pagar = ctx.pagar
    receber = ctx.receber
    inadimplencia = ctx.inadimplencia
    ref = Reference

    graficos = []
    graficos.append(_grafico_col(ind, 2, 1, ti, tf, "Faturamento por mês", E.NAVY))
    graficos.append(_grafico_col(ind, 9, 1, ti, tf, "Lucro líquido por mês", E.BLUE))
    graficos.append(_grafico_multiplo(ind, (11, 12), 1, ti, tf, "Entradas x saídas de caixa por mês",
                                      (E.NAVY, E.BLUE)))
    graficos.append(_grafico_linha(ind, 13, 1, ti, tf, "Evolução do saldo de caixa", E.BLUE))
    graficos.append(_grafico_linha(ind, 8, 1, ti, tf, "Margem de contribuição (%)", E.NAVY,
                                   fmt='0.0%'))
    graficos.append(_grafico_multiplo(ind, (14, 15), 1, ti, tf, "Orçado x realizado por mês",
                                      (E.BLUE, E.NAVY)))
    graficos.append(_grafico_col(dre, 3, 1, dc["corte"], dc["r1"],
                                 "Despesas por categoria (ano)", E.NAVY, altura=8.6, header=False))
    graficos.append(_grafico_col(dre, 3, 1, dc["r0"], dc["corte"] - 1,
                                 "Receitas por categoria (ano)", E.BLUE, altura=8.6, header=False))
    graficos.append(_grafico_col(pagar, 2, 1, apf[0], apf[1],
                                 "Contas a pagar por faixa de vencimento", E.NAVY))
    graficos.append(_grafico_col(receber, 2, 1, arf[0], arf[1],
                                 "Contas a receber por faixa de vencimento", E.BLUE))
    graficos.append(_grafico_col(inadimplencia, 2, 1, inad[0], inad[1],
                                 "Inadimplência por faixa de atraso", E.NAVY, altura=8.6))
    graficos.append(_grafico_col(emp, 17, 2, er[0], er[1], "Saldo devedor por contrato", E.BLUE,
                                 altura=8.6))
    linhas_gr = [34, 49, 64, 79, 94, 109]
    for k, ch in enumerate(graficos):
        linha = linhas_gr[k // 2]
        col = "B" if k % 2 == 0 else "J"
        ws.add_chart(ch, f"{col}{linha}")
    ctx.graficos_dashboard = graficos

    # ------------------------------------------------------- área de apoio
    apoio = 129
    U.mesclar(ws, f"B{apoio-1}:Q{apoio-1}",
              "ÁREA DE APOIO — CÁLCULOS DO DASHBOARD (não editar; os cartões acima referenciam estas "
              "células)", "SECAO", )
    ws[f"B{apoio-1}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    U.cabecalho(ws, apoio, 2, ["Valor calculado", "Linha calculada (base dos cartões acima)"], "TH2", 20)
    U.mesclar(ws, f"C{apoio}:Q{apoio}", "Linha calculada (base dos cartões acima)", "TH2")

    def filtro(col_lanc, cel):
        return f'IF({cel}="Todos","*",{cel})'

    def sumifs_dre(linha, tipo_extra=None):
        p = [f'{L_LINHA},"{linha}"', f'{L_ANO},$B$7', f'{L_MES},IF($D$7=0,">0",$D$7)',
             f'{L_CC},{filtro(L_CC, "$F$7")}', f'{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")}',
             f'{L_BANCO},{filtro(L_BANCO, "$J$7")}', f'{L_TIPO},{filtro(L_TIPO, "$L$7")}']
        return f'SUMIFS({L_VDRE},{",".join(p)})'

    def sumifs_caixa(tipo, col):
        ini = f'DATE($B$7,IF($D$7=0,1,$D$7),1)'
        fim = f'IF($D$7=0,DATE($B$7,12,31),EOMONTH(DATE($B$7,$D$7,1),0))'
        p = [f'{L_PAGTO},">="&{ini}', f'{L_PAGTO},"<="&{fim}', f'{L_CC},{filtro(L_CC, "$F$7")}',
             f'{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")}', f'{L_BANCO},{filtro(L_BANCO, "$J$7")}']
        return f'SUMIFS({col},{L_TIPO},"{tipo}",{",".join(p)})'

    # linha 131..145: cálculos do DRE filtrado
    linhas_calc = [
        (131, "Faturamento do período", f'={sumifs_dre("Faturamento Bruto")}'),
        (132, "Impostos e devoluções", f'={sumifs_dre("Impostos sobre Vendas")}+'
                                      f'{sumifs_dre("Devoluções")}'),
        (133, "Custos variáveis", f'={sumifs_dre("Custos Variáveis")}'),
        (134, "Despesas fixas", f'={sumifs_dre("Despesas Fixas")}'),
        (135, "Despesas financeiras", f'={sumifs_dre("Despesas Financeiras")}'),
        (136, "Receita líquida", "=$B$131-$B$132"),
        (137, "Receitas financeiras e outras", f'={sumifs_dre("Receitas Financeiras")}+'
                                               f'{sumifs_dre("Outras Receitas Operacionais")}'),
        (138, "Impostos sobre o lucro", f'={sumifs_dre("Impostos sobre o Resultado")}'),
        (139, "Margem de contribuição %", '=IFERROR(($B$136-$B$133)/$B$136,"")'),
        (140, "Lucro líquido", "=$B$136-$B$133-$B$134-$B$135+$B$137-$B$138"),
        (141, "Margem líquida %", '=IFERROR($B$140/$B$136,"")'),
        (142, "Entradas de caixa", f'={sumifs_caixa("Entrada", L_VFINAL)}'),
        (143, "Saídas de caixa", f'={sumifs_caixa("Saída", L_VFINAL)}'),
        (144, "Geração de caixa", "=$B$142-$B$143"),
        (145, "Margem de contribuição (R$)", "=$B$136-$B$133"),
    ]
    for linha, rot, f in linhas_calc:
        put(ws, linha, 2, f, "CALC_PCT" if "Margem" in rot else "CALC_MONEY")
        U.mesclar(ws, f"C{linha}:Q{linha}", rot, "CALC_TXT_SB")
    U.mesclar(ws, "C146:Q146", "Lançamentos no filtro", "CALC_TXT_SB")
    put(ws, 146, 2,
        f'=COUNTIFS({L_ANO},$B$7,{L_MES},IF($D$7=0,">0",$D$7),{L_CC},{filtro(L_CC, "$F$7")},'
        f'{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")},{L_BANCO},{filtro(L_BANCO, "$J$7")},'
        f'{L_TIPO},{filtro(L_TIPO, "$L$7")})', "CALC_INT")
    U.mesclar(ws, "C147:Q147", "A receber em aberto", "CALC_TXT_SB")
    put(ws, 147, 2,
        f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")},'
        f'{L_BANCO},{filtro(L_BANCO, "$J$7")},{L_CC},{filtro(L_CC, "$F$7")})', "CALC_MONEY")
    U.mesclar(ws, "C148:Q148", "A pagar em aberto", "CALC_TXT_SB")
    put(ws, 148, 2,
        f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")},'
        f'{L_BANCO},{filtro(L_BANCO, "$J$7")},{L_CC},{filtro(L_CC, "$F$7")})', "CALC_MONEY")
    U.mesclar(ws, "C149:Q149", "Vencido a receber (títulos em atraso)", "CALC_TXT_SB")
    put(ws, 149, 2,
        f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_SALDO},">0",{L_VENC},"<"&{HOJE},'
        f'{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")})', "CALC_MONEY")
    U.mesclar(ws, "C150:Q150", "Vencido a pagar (títulos em atraso)", "CALC_TXT_SB")
    put(ws, 150, 2,
        f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},"<"&{HOJE},'
        f'{L_SUBCAT},{filtro(L_SUBCAT, "$H$7")})', "CALC_MONEY")
    U.mesclar(ws, "C151:Q151", "Orçado no período", "CALC_TXT_SB")
    r_orc, r_real = ctx.orc_rows
    r_orc_hdr = r_orc - 1
    put(ws, 151, 2,
        f'=IF($D$7=0,IFERROR(SUM(\'{P.ABA_ORC}\'!$O$7:$O${P.ORC_FIM}),0),'
        f'SUMIF(\'{P.ABA_ORC}\'!$B${r_orc_hdr}:$M${r_orc_hdr},$D$7,'
        f'\'{P.ABA_ORC}\'!$B${r_orc}:$M${r_orc}))', "CALC_MONEY")
    U.mesclar(ws, "C152:Q152", "Realizado no período", "CALC_TXT_SB")
    put(ws, 152, 2,
        f'=IF($D$7=0,IFERROR(SUM(\'{P.ABA_ORC}\'!$AB$7:$AB${P.ORC_FIM}),0),'
        f'SUMIF(\'{P.ABA_ORC}\'!$B${r_orc_hdr}:$M${r_orc_hdr},$D$7,'
        f'\'{P.ABA_ORC}\'!$B${r_real}:$M${r_real}))', "CALC_MONEY")
    U.mesclar(ws, "C153:Q153", "Consumo do orçamento (%)", "CALC_TXT_SB")
    put(ws, 153, 2, '=IFERROR($B$152/$B$151,"")', "CALC_PCT")

    U.mesclar(ws, "B154:Q154", "Os cartões usam exatamente estas células. Filtros: ano (B7), mês (D7, "
                               "0 = ano todo), centro de custo (F7), categoria (H7), conta (J7) e tipo "
                               "(L7) — use 'Todos' para não filtrar.", "NOTA")
    ws.row_dimensions[apoio - 1].height = 20
    ws.row_dimensions[apoio].height = 30
    U.cf_barras(ws, "B131:B145", E.BLUE)

    U.preparar(ws, grade=False, zoom=85, tab_cor=E.NAVY)
    ws.sheet_view.zoomScale = 85
    ctx.dashboard = ws
    return ws


# ==========================================================================
# PAINEL DE CONTROLE
# ==========================================================================
def painel_controle(wb, ctx):
    ws = wb.create_sheet(P.ABA_PAINEL)
    ws.sheet_view.showGridLines = False
    U.larguras(ws, {"A": 3, "B": 30, "C": 30, "D": 30, "E": 30, "F": 3, "G": 3})
    U.mesclar(ws, "B2:E3", f'=IF({FANTASIA}="","GESTÃO FINANCEIRA EMPRESARIAL",{FANTASIA})', "TITULO_DASH")
    ws.row_dimensions[2].height = 30
    U.mesclar(ws, "B4:E4",
              "Sistema completo de gestão financeira: lançamentos, contas a pagar e a receber, fluxo de "
              "caixa, DRE gerencial, orçamento, empréstimos, indicadores e dashboard executivo.", "SUBTITULO_DASH")
    U.mesclar(ws, "B5:E5", f'="Hoje: "&TEXT({HOJE},"DD/MM/YYYY")&"     |     Ano de análise: "&'
                          f'{ANO}&"     |     Mês de análise: "&IF({MES}=0,"ano completo",'
                          f'TEXT(DATE({ANO},{MES},1),"MM/YYYY"))&"     |     Arquivo: '
                          f'Excel (xlsx) e Google Sheets"', "NOTA_WRAP")

    # ---------------------------------------------------------------- indicadores
    U.secao(ws, 7, "SITUAÇÃO ATUAL", span=4, col=2, altura=22)
    cards = [
        ("Saldo de caixa disponível",
         f'={DISPONIVEL}+SUMIFS({L_VFINAL},{L_TIPO},"Entrada",{L_PAGTO},"<="&{HOJE})'
         f'-SUMIFS({L_VFINAL},{L_TIPO},"Saída",{L_PAGTO},"<="&{HOJE})', "KPI_VALUE"),
        ("Faturamento do mês de análise",
         f'=SUMIFS({L_VDRE},{L_LINHA},"Faturamento Bruto",{L_ANO},{ANO},{L_MES},{MES})', "KPI_VALUE"),
        ("A receber em aberto", f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada")', "KPI_VALUE"),
        ("A pagar em aberto", f'=SUMIFS({L_SALDO},{L_TIPO},"Saída")', "KPI_VALUE"),
    ]
    for i, (rot, val, est) in enumerate(cards):
        U.cartao(ws, 8 + (i // 2) * 3, 2 + (i % 2) * 2, 2, rot, val, est)

    # ---------------------------------------------------------------- botões
    U.secao(ws, 15, "MENU — CLIQUE NAS TELAS ABAIXO PARA NAVEGAR", span=4, col=2, altura=22)
    botoes = [
        ("📊 DASHBOARD", "Visão executiva: cartões, gráficos e alertas.", P.ABA_DASH),
        ("✍️ LANÇAR ENTRADA", "Registre uma receita na aba LANÇAMENTOS.", P.ABA_LANC),
        ("✍️ LANÇAR SAÍDA", "Registre uma despesa na aba LANÇAMENTOS.", P.ABA_LANC),
        ("📄 CONTAS A PAGAR", "Títulos a pagar, aging e prioridades.", P.ABA_AP),
        ("📥 CONTAS A RECEBER", "Carteira, recebimentos e prazos.", P.ABA_AR),
        ("💧 FLUXO DE CAIXA", "Entradas, saídas, saldos e projeções.", P.ABA_FLUXO),
        ("📈 DRE GERENCIAL", "Resultado por competência e margens.", P.ABA_DRE),
        ("👥 CLIENTES", "Cadastro e indicadores por cliente.", P.ABA_CLI),
        ("🏭 FORNECEDORES", "Cadastro e indicadores por fornecedor.", P.ABA_FORN),
        ("🏦 BANCOS E CARTÕES", "Contas, saldos, limites e faturas.", P.ABA_BANCOS),
        ("🧾 RELATÓRIOS", "Relatórios resumidos para imprimir.", P.ABA_REL),
        ("⚙️ CONFIGURAÇÕES", "Empresa, metas, parâmetros e listas.", P.ABA_CONFIG),
    ]
    for i, (rotulo, desc, aba) in enumerate(botoes):
        r = 16 + (i // 2) * 3
        col = 2 + (i % 2) * 2
        rng = f"{L(col)}{r}:{L(col+1)}{r}"
        c = U.mesclar(ws, rng, rotulo, "MENU_BOTAO")
        U.mesclar(ws, f"{L(col)}{r+1}:{L(col+1)}{r+1}", desc, "MENU_BOTAO_D")
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{aba}'!A1", display=rotulo,
                                tooltip=f"Abrir {aba}")
    U.mesclar(ws, "B34:E34", "Dica: os lançamentos são a base de tudo — cadastre cliente/fornecedor e "
                             "categoria antes de lançar. Todas as demais abas se atualizam sozinhas.",
              "NOTA_WRAP")
    ws.row_dimensions[34].height = 26

    # ---------------------------------------------------------------- passos
    U.secao(ws, 36, "COMECE POR AQUI — ROTEIRO DE 6 PASSOS", span=4, col=2, altura=22)
    passos = [
        ("1. Configurações", "Informe os dados da empresa, as metas e o ano/mês de análise. Confira as "
                             "listas auxiliares que alimentam os menus suspensos.", P.ABA_CONFIG),
        ("2. Cadastros", "Cadastre clientes, fornecedores, plano de contas, centros de custo, contas "
                         "bancárias e cartões.", P.ABA_CLI),
        ("3. Lançamentos", "Registre todas as entradas e saídas. Use os menus suspensos e preencha as "
                           "células amarelas; as colunas cinzas são calculadas.", P.ABA_LANC),
        ("4. Conferência", "Verifique contas a pagar e a receber, a conciliação bancária e a coluna de "
                           "validação dos lançamentos.", P.ABA_AP),
        ("5. Análise", "Acompanhe o dashboard, o DRE, o fluxo de caixa, o orçamento e os indicadores.",
         P.ABA_DASH),
        ("6. Planejamento", "Use ponto de equilíbrio, projeção de caixa, simulador de cenários e despesas "
                            "recorrentes para decidir.", P.ABA_PE),
    ]
    for i, (titulo, texto, aba) in enumerate(passos):
        r = 37 + i
        c = U.mesclar(ws, f"B{r}:B{r}", titulo, "KPI_LABEL")
        U.mesclar(ws, f"C{r}:E{r}", texto, "CALC_TXT")
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{aba}'!A1", display=titulo,
                                tooltip=f"Abrir {aba}")
        ws.row_dimensions[r].height = 18
    U.mesclar(ws, "B44:E44",
              "⚠ Esta planilha acompanha DADOS DE EXEMPLO (empresa fictícia) para você testar todos os "
              "recursos. Antes de usar de verdade, apague os dados de exemplo — veja o passo a passo na "
              "aba INSTRUÇÕES.", "ALERTA_TXT")
    ws.row_dimensions[44].height = 30
    U.mesclar(ws, "B45:E45", "Versão 1.0 — 25 abas — Compatível com Microsoft Excel 2016+ e Google Sheets.",
              "NOTA_WRAP")
    U.preparar(ws, grade=False, tab_cor=E.NAVY)
    ctx.painel = ws
    return ws


# ==========================================================================
# RELATÓRIOS
# ==========================================================================
def relatorios(wb, ctx):
    ws = wb.create_sheet(P.ABA_REL)
    ws.sheet_view.showGridLines = False
    U.larguras(ws, {"A": 3, "B": 30, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16,
                    "J": 16, "K": 16, "L": 16, "M": 16, "N": 16, "O": 14})
    U.mesclar(ws, "B2:O2", "CENTRAL DE RELATÓRIOS — RESUMOS PRONTOS PARA IMPRIMIR", "TITULO_DASH")
    ws.row_dimensions[2].height = 26
    U.mesclar(ws, "B3:O3",
              f'=IF({RAZAO}="","",{RAZAO}&"  |  CNPJ: ")&'
              f'IF({RAZAO}="","",{P.CONFIG_REF["cnpj"]}&"  |  Emitido em "&TEXT({HOJE},"DD/MM/YYYY"))',
              "SUBTITULO_DASH")
    ti, tf = ctx.ind_trend
    dc = ctx.dre_cat
    dre = ctx.dre

    # ----------------------------------------------------------- índice
    U.secao(ws, 5, "1. ÍNDICE DOS RELATÓRIOS DISPONÍVEIS", span=8, col=2, altura=22)
    U.cabecalho(ws, 6, 2, ["Nº", "Relatório", "Onde encontrar", "O que apresenta"], "TH2", 30)
    indice = [
        ("1", "Demonstrativo de resultado (DRE gerencial)", "DRE GERENCIAL",
         "Receita, custos, despesas, margens e lucro mês a mês."),
        ("2", "Fluxo de caixa realizado e projetado", "FLUXO DE CAIXA",
         "Entradas, saídas, resultado, saldo inicial e final, por mês."),
        ("3", "Contas a pagar por vencimento", "CONTAS A PAGAR",
         "Vencidos, vence hoje, 7, 30 dias, aging e prioridades."),
        ("4", "Contas a receber e carteira", "CONTAS A RECEBER",
         "Carteira por faixa, recebimentos, prazos e ticket médio."),
        ("5", "Inadimplência por faixa de atraso", "INADIMPLÊNCIA",
         "Valores vencidos, aging 1-7 a +90 dias e ranking de clientes."),
        ("6", "Orçamento x realizado", "ORÇAMENTO X REALIZADO",
         "Orçado, realizado, diferença, % consumido e alertas."),
        ("7", "Endividamento e serviço da dívida", "EMPRÉSTIMOS E FINANCIAMENTOS",
         "Contratos, parcelas, juros, saldo devedor e comprometimento."),
        ("8", "Resultado por centro de custo", "CENTRO DE CUSTOS",
         "Receita, custos, despesas, resultado e margem por centro."),
        ("9", "Indicadores de desempenho", "INDICADORES",
         "KPIs de vendas, rentabilidade, caixa, prazos e dívida."),
        ("10", "Conciliação bancária", "CONCILIAÇÃO BANCÁRIA",
         "Extrato x planilha, pendências e divergências."),
        ("11", "Projeção de caixa (7 a 180 dias)", "PROJEÇÃO DE CAIXA",
         "Saldo projetado por horizonte e alerta de déficit."),
        ("12", "Comparativo de cenários", "SIMULADOR DE CENÁRIOS",
         "Cenários atual, otimista e conservador com lucro e margem."),
        ("13", "Despesas recorrentes", "DESPESAS RECORRENTES",
         "Contratos recorrentes, próximas ocorrências e projeção."),
        ("14", "Dashboard executivo", "DASHBOARD",
         "16 cartões, 18 alertas e 12 gráficos gerenciais."),
    ]
    r0 = 7
    for i, (n, rel, onde, oque) in enumerate(indice):
        r = r0 + i
        put(ws, r, 2, f'="{n}"', "CALC_TXT_C")
        put(ws, r, 3, rel, "CALC_TXT")
        put(ws, r, 4, onde, "LINK_ABA")
        ws.cell(row=r, column=4).hyperlink = f"#'{onde}'!A1"
        U.mesclar(ws, f"E{r}:I{r}", oque, "NOTA")
    r_idx_fim = r0 + len(indice) - 1

    # ------------------------------------------------- resumo mensal
    def bloco(r, titulo, cab, dados, largura_nota=None):
        U.secao(ws, r, titulo, span=8, col=2, altura=22)
        U.cabecalho(ws, r + 1, 2, cab, "TH2", 28)
        return r + 2

    r = r_idx_fim + 2
    r = bloco(r, "2. FLUXO DE CAIXA — RESUMO MENSAL DO ANO DE ANÁLISE",
              ["Mês", "Entradas realizadas", "Saídas realizadas", "Resultado", "Saldo final",
               "Faturamento (competência)", "Lucro líquido", "Margem líquida %"], [])
    r2 = r
    for k in range(12):
        rr = r2 + k
        src = ti + k
        put(ws, rr, 2, f"='{P.ABA_IND}'!$A${src}", "CALC_DATE")
        put(ws, rr, 3, f"='{P.ABA_IND}'!$K${src}", "CALC_MONEY")
        put(ws, rr, 4, f"='{P.ABA_IND}'!$L${src}", "CALC_MONEY")
        put(ws, rr, 5, f"=$C{rr}-$D{rr}", "CALC_MONEY")
        put(ws, rr, 6, f"='{P.ABA_IND}'!$M${src}", "CALC_MONEY")
        put(ws, rr, 7, f"='{P.ABA_IND}'!$B${src}", "CALC_MONEY")
        put(ws, rr, 8, f"='{P.ABA_IND}'!$I${src}", "CALC_MONEY")
        put(ws, rr, 9, f"='{P.ABA_IND}'!$J${src}", "CALC_PCT")
    r2_fim = r2 + 11
    put(ws, r2_fim + 1, 2, "TOTAL", "TOTAL_TXT")
    for col in (3, 4, 5, 7, 8):
        put(ws, r2_fim + 1, col, f"=SUM({L(col)}{r2}:{L(col)}{r2_fim})", "TOTAL_MONEY")
    put(ws, r2_fim + 1, 6, f"=$F{r2_fim}", "TOTAL_MONEY")
    put(ws, r2_fim + 1, 9, f'=IF($G{r2_fim+1}=0,"",$H{r2_fim+1}/$G{r2_fim+1})', "TOTAL_PCT")
    U.cf_barras(ws, f"C{r2}:C{r2_fim}")

    # ------------------------------------------------- DRE resumida
    r = r2_fim + 3
    r = bloco(r, "3. DRE GERENCIAL — RESUMO DO ANO DE ANÁLISE",
              ["Linha do DRE", "Janeiro a dezembro", "% da receita líquida"], [])
    r3 = r
    dre_linhas = [
        ("Faturamento bruto", "=$B$131" if False else f"=SUMIFS({L_VDRE},{L_LINHA},\"Faturamento Bruto\",{L_ANO},{ANO})", "CALC_MONEY"),
        ("(-) Impostos sobre vendas", f'=SUMIFS({L_VDRE},{L_LINHA},"Impostos sobre Vendas",{L_ANO},{ANO})', "CALC_MONEY"),
        ("(-) Devoluções", f'=SUMIFS({L_VDRE},{L_LINHA},"Devoluções",{L_ANO},{ANO})', "CALC_MONEY"),
        ("▶ Receita líquida", "=$C$%d-$C$%d-$C$%d" % (r3, r3 + 1, r3 + 2), "TOTAL_MONEY"),
        ("(-) Custos variáveis", f'=SUMIFS({L_VDRE},{L_LINHA},"Custos Variáveis",{L_ANO},{ANO})', "CALC_MONEY"),
        ("▶ Margem de contribuição", "=$C$%d-$C$%d" % (r3 + 3, r3 + 4), "TOTAL_MONEY"),
        ("(-) Despesas fixas", f'=SUMIFS({L_VDRE},{L_LINHA},"Despesas Fixas",{L_ANO},{ANO})', "CALC_MONEY"),
        ("▶ EBITDA / resultado operacional", "=$C$%d-$C$%d" % (r3 + 5, r3 + 6), "TOTAL_MONEY"),
        ("(-) Despesas financeiras", f'=SUMIFS({L_VDRE},{L_LINHA},"Despesas Financeiras",{L_ANO},{ANO})', "CALC_MONEY"),
        ("(+) Receitas financeiras", f'=SUMIFS({L_VDRE},{L_LINHA},"Receitas Financeiras",{L_ANO},{ANO})', "CALC_MONEY"),
        ("(+) Outras receitas", f'=SUMIFS({L_VDRE},{L_LINHA},"Outras Receitas Operacionais",{L_ANO},{ANO})', "CALC_MONEY"),
        ("▶ Lucro antes dos impostos", "=$C$%d-$C$%d+$C$%d+$C$%d" % (r3 + 7, r3 + 8, r3 + 9, r3 + 10),
         "TOTAL_MONEY"),
        ("(-) Impostos sobre o resultado", f'=SUMIFS({L_VDRE},{L_LINHA},"Impostos sobre o Resultado",{L_ANO},{ANO})', "CALC_MONEY"),
        ("▶ Lucro líquido", "=$C$%d-$C$%d" % (r3 + 11, r3 + 12), "TOTAL_MONEY"),
        ("(-) Distribuição de lucros", f'=SUMIFS({L_VDRE},{L_LINHA},"Distribuição de Lucros",{L_ANO},{ANO})', "CALC_MONEY"),
        ("▶ Lucro retido no ano", "=$C$%d-$C$%d" % (r3 + 13, r3 + 14), "TOTAL_MONEY"),
    ]
    for i, (rot, f, est) in enumerate(dre_linhas):
        rr = r3 + i
        put(ws, rr, 2, rot, "CALC_TXT")
        put(ws, rr, 3, f, est)
        put(ws, rr, 4, f'=IF($C${r3+3}=0,"",{-1 if "(-)" in rot else 1}*$C{rr}/$C${r3+3})', "CALC_PCT")
    r3_fim = r3 + len(dre_linhas) - 1
    U.cf_barras(ws, f"C{r3}:C{r3_fim}")

    # ---------------------------------------------------- contas a pagar/receber
    r = r3_fim + 3
    r = bloco(r, "4. CONTAS A PAGAR E A RECEBER — POSIÇÃO NA DATA DE HOJE",
              ["Indicador", "Valor", "Quantidade", "Observação"], [])
    r4 = r
    pos = [
        ("Total a pagar em aberto", f'=SUMIFS({L_SALDO},{L_TIPO},"Saída")',
         f'=COUNTIFS({L_TIPO},"Saída",{L_SALDO},">0")', "Compromissos futuros registrados."),
        ("Total vencido a pagar", f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},"<"&{HOJE})',
         f'=COUNTIFS({L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},"<"&{HOJE})', "Exige ação imediata."),
        ("A pagar em 7 dias",
         f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+7)',
         f'=COUNTIFS({L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+7)',
         "Programe os pagamentos da semana."),
        ("A pagar em 30 dias",
         f'=SUMIFS({L_SALDO},{L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+30)',
         f'=COUNTIFS({L_TIPO},"Saída",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+30)',
         "Planejamento do mês."),
        ("Total a receber em aberto", f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_LINHA},"Faturamento Bruto")',
         f'=COUNTIFS({L_TIPO},"Entrada",{L_LINHA},"Faturamento Bruto",{L_SALDO},">0")', "Carteira de recebíveis."),
        ("Total vencido a receber",
         f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_LINHA},"Faturamento Bruto",{L_SALDO},">0",{L_VENC},"<"&{HOJE})',
         f'=COUNTIFS({L_TIPO},"Entrada",{L_LINHA},"Faturamento Bruto",{L_SALDO},">0",{L_VENC},"<"&{HOJE})',
         "Base da régua de cobrança."),
        ("A receber em 7 dias",
         f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+7)',
         f'=COUNTIFS({L_TIPO},"Entrada",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+7)',
         "Confirme com os clientes."),
        ("A receber em 30 dias",
         f'=SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+30)',
         f'=COUNTIFS({L_TIPO},"Entrada",{L_SALDO},">0",{L_VENC},">="&{HOJE},{L_VENC},"<="&{HOJE}+30)',
         "Previsão de entradas do mês."),
        ("Índice de inadimplência",
         f'=IFERROR(SUMIFS({L_SALDO},{L_TIPO},"Entrada",{L_LINHA},"Faturamento Bruto",{L_SALDO},">0",'
         f'{L_VENC},"<"&{HOJE})/SUMIFS({L_VPREV},{L_TIPO},"Entrada",{L_LINHA},"Faturamento Bruto",'
         f'{L_ANO},{ANO}),"")', '=""', "Compare com o limite definido nas configurações."),
    ]
    for i, (rot, val, qtd, obs) in enumerate(pos):
        rr = r4 + i
        put(ws, rr, 2, rot, "CALC_TXT")
        est = "CALC_PCT" if "inadimpl" in rot.lower() else "CALC_MONEY"
        put(ws, rr, 3, val, est)
        put(ws, rr, 4, qtd, "CALC_INT")
        U.mesclar(ws, f"E{rr}:I{rr}", obs, "NOTA")
    U.cf_maior(ws, f"C{r4+1}:C{r4+1}", 0, E.RED_BG, E.RED_TX)
    U.cf_maior(ws, f"C{r4+5}:C{r4+5}", 0, E.RED_BG, E.RED_TX)

    # ------------------------------------------------- orçamento x realizado
    r = r4 + len(pos) + 2
    r = bloco(r, "5. ORÇAMENTO X REALIZADO — RESUMO MENSAL",
              ["Mês", "Orçado", "Realizado", "Diferença", "% consumido", "Situação"], [])
    r5 = r
    orc = ctx.orcamento
    for k in range(12):
        rr = r5 + k
        put(ws, rr, 2, f'=TEXT(DATE({ANO},{k+1},1),"MM/AAAA")', "CALC_TXT_C")
        put(ws, rr, 3, f"='{P.ABA_ORC}'!{L(2+k)}{ctx.orc_rows[0]}", "CALC_MONEY")
        put(ws, rr, 4, f"='{P.ABA_ORC}'!{L(2+k)}{ctx.orc_rows[1]}", "CALC_MONEY")
        put(ws, rr, 5, f"=$C{rr}-$D{rr}", "CALC_MONEY")
        put(ws, rr, 6, f'=IF($C{rr}=0,"",$D{rr}/$C{rr})', "CALC_PCT")
        put(ws, rr, 7, f'=IF($C{rr}=0,"Sem orçamento",IF($D{rr}>$C{rr},"🔴 Estourou",'
                       f'IF($D{rr}>$C{rr}*0.9,"🟡 > 90%","🟢 Dentro")))', "CALC_TXT_C")
    r5_fim = r5 + 11
    put(ws, r5_fim + 1, 2, "TOTAL DO ANO", "TOTAL_TXT")
    for col in (3, 4, 5):
        put(ws, r5_fim + 1, col, f"=SUM({L(col)}{r5}:{L(col)}{r5_fim})", "TOTAL_MONEY")
    put(ws, r5_fim + 1, 6, f'=IF($C{r5_fim+1}=0,"",$D{r5_fim+1}/$C{r5_fim+1})', "TOTAL_PCT")
    U.cf_contem(ws, f"G{r5}:G{r5_fim}", "🔴", E.RED_BG, E.RED_TX)
    U.cf_contem(ws, f"G{r5}:G{r5_fim}", "🟡", E.YEL_BG, E.YEL_TX)
    U.cf_contem(ws, f"G{r5}:G{r5_fim}", "🟢", E.GREEN_BG, E.GREEN_TX)

    # ----------------------------------------------------------- por centro de custo
    r = r5_fim + 3
    r = bloco(r, "6. RESULTADO POR CENTRO DE CUSTO (ano de análise)",
              ["Centro de custo", "Receitas", "Custos e despesas", "Resultado", "Margem %"], [])
    r6 = r
    for i in range(len(D.CENTROS)):
        rr = r6 + i
        put(ws, rr, 2, f"='{P.ABA_CC}'!$B${P.CC_INI + i}", "CALC_TXT")
        put(ws, rr, 3, f'=SUMIFS({L_VDRE},{L_CC},$B{rr},{L_TIPO},"Entrada",{L_ANO},{ANO})', "CALC_MONEY")
        put(ws, rr, 4, f'=SUMIFS({L_VDRE},{L_CC},$B{rr},{L_TIPO},"Saída",{L_ANO},{ANO})', "CALC_MONEY")
        put(ws, rr, 5, f"=$C{rr}-$D{rr}", "CALC_MONEY")
        put(ws, rr, 6, f'=IF($C{rr}=0,"",$E{rr}/$C{rr})', "CALC_PCT")
    r6_fim = r6 + len(D.CENTROS) - 1
    U.cf_barras(ws, f"C{r6}:C{r6_fim}")

    U.mesclar(ws, f"B{r6_fim+2}:O{r6_fim+3}",
              "IMPRESSÃO: cada bloco acima pode ser impresso separadamente. Use Visualizar impressão e "
              "ajuste a área de impressão para o bloco desejado (layout paisagem, ajustar à largura).",
              "NOTA_WRAP")
    U.preparar(ws, grade=False, tab_cor="7030A0", area_impressao=f"A1:I{r6_fim+4}")
    ctx.relatorios = ws
    return ws


# ==========================================================================
# INSTRUÇÕES
# ==========================================================================
def instrucoes(wb, ctx):
    ws = wb.create_sheet(P.ABA_INSTR)
    ws.sheet_view.showGridLines = False
    U.larguras(ws, {"A": 3, "B": 4, "C": 34, "D": 30, "E": 34, "F": 30, "G": 4})
    U.mesclar(ws, "B2:F2", "INSTRUÇÕES DE USO — SISTEMA DE GESTÃO FINANCEIRA EMPRESARIAL", "TITULO_DASH")
    ws.row_dimensions[2].height = 28
    U.mesclar(ws, "B3:F3",
              "Este arquivo funciona no Microsoft Excel (2016 ou superior) e no Google Sheets. "
              "Foi construído com fórmulas nativas (SUMIFS, COUNTIFS, IFERROR, INDEX/MATCH), "
              "validação de dados e formatação condicional.", "SUBTITULO_DASH")

    linhas = [
        ("SEÇÃO", "1. O QUE ESTE SISTEMA FAZ", None),
        ("TEXTO", "É um sistema completo de gestão financeira para pequenas e médias empresas: controla "
                  "lançamentos, contas a pagar e a receber, fluxo de caixa, resultado (DRE), orçamento, "
                  "empréstimos, conciliação bancária, inadimplência, indicadores e um dashboard executivo "
                  "com gráficos e alertas automáticos.", None),
        ("TEXTO", "Todos os números são calculados automaticamente a partir da aba LANÇAMENTOS. Você só "
                  "preenche as células amarelas; as células cinzas contêm as fórmulas.", None),
        ("SEÇÃO", "2. ESTRUTURA DAS ABAS", None),
        ("TABELA", "PAINEL DE CONTROLE", "Menu inicial com atalhos para todas as telas e a situação atual.", None),
        ("TABELA", "DASHBOARD", "16 cartões, 18 alertas automáticos e 12 gráficos com 6 filtros.", None),
        ("TABELA", "CONFIGURAÇÕES", "Dados da empresa, ano/mês de análise, metas e listas auxiliares.", None),
        ("TABELA", "LANÇAMENTOS", "Base única de entradas e saídas (competência, vencimento, pagamento).", None),
        ("TABELA", "CONTAS A PAGAR / A RECEBER", "Classificação automática por vencimento, aging e prazos.", None),
        ("TABELA", "FLUXO DE CAIXA", "Diário, mensal e por tipo de operação (operacional/invest./financ.).", None),
        ("TABELA", "DRE GERENCIAL", "Resultado por competência, margens e comparação mensal.", None),
        ("TABELA", "ORÇAMENTO X REALIZADO", "Orçado por categoria e mês, realizado, % e alertas.", None),
        ("TABELA", "EMPRÉSTIMOS E FINANCIAMENTOS", "Contratos, cronograma SAC/PRICE, juros e saldo devedor.", None),
        ("TABELA", "CONCILIAÇÃO BANCÁRIA", "Extrato x planilha com pendências e divergências.", None),
        ("TABELA", "INADIMPLÊNCIA", "Faixas de atraso (1-7 a +90 dias) e ranking de clientes.", None),
        ("TABELA", "INDICADORES", "KPIs por área e a área ATENÇÃO DO GESTOR com 18 verificações.", None),
        ("TABELA", "PONTO DE EQUILÍBRIO", "Faturamento mínimo, margem de segurança e meta de lucro.", None),
        ("TABELA", "SIMULADOR DE CENÁRIOS", "Compara cenários atual, otimista e conservador.", None),
        ("TABELA", "PROJEÇÃO DE CAIXA", "Saldo projetado em 7, 15, 30, 60, 90 e 180 dias.", None),
        ("TABELA", "DESPESAS RECORRENTES", "Cadastro das despesas fixas e geração de lançamentos futuros.", None),
        ("TABELA", "RELATÓRIOS", "14 relatórios resumidos, prontos para imprimir.", None),
        ("SEÇÃO", "3. LEGENDA DE CORES E TIPOS DE CÉLULA", None),
        ("COR", "Amarelo (FFF2CC)", "Célula de entrada: você pode e deve editar.", "IN_TXT_C"),
        ("COR", "Cinza (F2F2F2)", "Célula calculada: contém fórmula, não edite.", "CALC_TXT_C"),
        ("COR", "Cabeçalho azul", "Títulos de tabelas e colunas.", "TH2"),
        ("COR", "Verde / Amarelo / Vermelho", "Status: dentro do esperado, atenção, alerta.", None),
        ("SEÇÃO", "4. COMO COMEÇAR (ROTEIRO)", None),
        ("PASSO", "1. Abra a aba CONFIGURAÇÕES", "Informe razão social, CNPJ, ano e mês de análise e as "
                                                 "metas (faturamento, lucro, margem, saldo mínimo, prazos).", None),
        ("PASSO", "2. Revise o PLANO DE CONTAS", "As categorias e subcategorias alimentam os menus "
                                                 "suspensos de todas as abas. Inclua as suas.", None),
        ("PASSO", "3. Cadastre clientes, fornecedores, centros de custo, contas e cartões", "Use as abas de "
                  "cadastro; os nomes devem ser idênticos aos usados nos lançamentos.", None),
        ("PASSO", "4. Lance as entradas e saídas na aba LANÇAMENTOS", "Comece pela data de competência, "
                  "tipo, categoria e subcategoria; informe previsto, realizado, vencimento e pagamento.", None),
        ("PASSO", "5. Confira as abas de análise", "Contas a pagar/receber, fluxo de caixa, DRE, "
                  "orçamento, indicadores e dashboard.", None),
        ("PASSO", "6. Use as ferramentas de decisão", "Ponto de equilíbrio, projeção de caixa, simulador "
                  "de cenários e despesas recorrentes.", None),
        ("SEÇÃO", "5. COMO LIMPAR OS DADOS DE EXEMPLO", None),
        ("TEXTO", "Esta planilha acompanha dados fictícios (empresa /Aurora Distribuidora/, 3 contas "
                  "bancárias, 3 empréstimos, cerca de 700 lançamentos, orçamento, conciliação e despesas "
                  "recorrentes) para você testar todos os recursos.", None),
        ("PASSO", "LANÇAMENTOS", "Selecione as células das colunas A até X (entrada) e pressione Delete. "
                                 "As colunas de cálculo (Y até AO) contêm fórmulas e podem permanecer.", None),
        ("PASSO", "CADASTROS", "Nas abas de cadastro, apague o conteúdo das linhas de dados mantendo as "
                               "fórmulas das colunas calculadas.", None),
        ("PASSO", "ORÇAMENTO E RECORRENTES", "Apague os valores das colunas de orçamento e as linhas das "
                                             "despesas recorrentes de exemplo.", None),
        ("SEÇÃO", "6. COMO USAR NO GOOGLE SHEETS", None),
        ("PASSO", "1. Importar", "No Google Drive, clique em Novo > Fazer upload de arquivo e selecione o "
                                 "arquivo .xlsx (ou abra-o e use Arquivo > Salvar como Planilhas Google).", None),
        ("PASSO", "2. Conferir", "As fórmulas, validações e formatação condicional são convertidas "
                                 "automaticamente. Confira as listas suspensas e os gráficos.", None),
        ("PASSO", "3. Cuidados", "Evite colar dados como texto em colunas de data/valor. Em caso de aviso "
                                 "de formatação, use Formatar > Número > Moeda/Data.", None),
        ("SEÇÃO", "7. ROTINA MENSAL RECOMENDADA", None),
        ("TEXTO", "1) Lance os títulos do mês (vendas, compras, despesas e tributos). 2) Atualize o "
                  "realizado e as datas de pagamento/recebimento. 3) Confira contas a pagar e a receber. "
                  "4) Faça a conciliação bancária. 5) Confira o DRE e os indicadores. 6) Compare orçado x "
                  "realizado. 7) Analise a projeção de caixa e a inadimplência. 8) Registre o orçamento do "
                  "mês seguinte e revise as despesas recorrentes.", None),
        ("SEÇÃO", "8. ALERTAS AUTOMÁTICOS (ATENÇÃO DO GESTOR)", None),
        ("TEXTO", "As 18 verificações cobrem: caixa abaixo do mínimo, lucro e margem abaixo da meta, "
                  "faturamento abaixo da meta, inadimplência acima do limite, títulos pagos/recebidos "
                  "vencidos, compromissos de 7 dias maiores que o saldo, projeção de 30 dias, orçamento "
                  "estourado, erros de preenchimento, clientes em atraso, prazos médios fora da meta, "
                  "endividamento elevado, despesas fixas crescendo e reserva de caixa. Confira na aba "
                  "INDICADORES e no DASHBOARD.", None),
        ("SEÇÃO", "9. PERGUNTAS FREQUENTES", None),
        ("FAQ", "Posso acrescentar linhas?", "Sim: a aba LANÇAMENTOS suporta até 2.000 lançamentos, 60 "
                                             "clientes, 60 fornecedores, 80 contas no plano, 500 linhas de "
                                             "conciliação, 10 contratos com 60 parcelas cada. Para ampliar, "
                                             "copie a última linha e cole para baixo.", None),
        ("FAQ", "Por que não consigo editar uma célula?", "As células calculadas estão protegidas. Para "
                                                          "desproteger: Revisão > Desproteger Planilha "
                                                          "(sem senha).", None),
        ("FAQ", "As fórmulas funcionam em português?", "As fórmulas usam funções nativas compatíveis "
                                                       "(SUMIFS, COUNTIFS, IFERROR, INDEX/MATCH, PMT). No "
                                                       "Excel em português elas aparecem traduzidas; no "
                                                       "Google Sheets também.", None),
        ("FAQ", "E se eu usar cartão de crédito?", "Cadastre o cartão na aba CONTAS BANCÁRIAS e CARTÕES. "
                                                   "Nos lançamentos, escolha a forma Cartão de crédito e a "
                                                   "conta correspondente.", None),
        ("FAQ", "Como funcionam as vendas parceladas?", "Informe o número da parcela e o total de parcelas. "
                                                        "Cada parcela é um lançamento com o mesmo cliente e "
                                                        "categoria; o valor previsto é o da parcela.", None),
        ("SEÇÃO", "10. VERIFICAÇÕES REALIZADAS (TESTES)", None),
        ("TEXTO", "A planilha foi testada com dados fictícios: conferência dos saldos bancários, do fluxo "
                  "de caixa, do DRE, do lucro e das margens, dos indicadores do dashboard, dos gráficos, "
                  "dos alertas e do orçamento x realizado. Nenhuma fórmula apresentou erro (#DIV/0!, "
                  "#N/A, #VALOR!, #REF!).", None),
        ("TEXTO", "As únicas mensagens possíveis são as de alerta intencional (por exemplo, '⚠ Cliente/"
                  "Fornecedor não cadastrado') quando algum dado de entrada estiver inconsistente.", None),
    ]
    r = 5
    for item in linhas:
        tipo, a = item[0], item[1]
        b = item[2] if len(item) > 2 else None
        estilo = item[3] if len(item) > 3 else None
        if tipo == "SEÇÃO":
            U.secao(ws, r, a, span=4, col=2, altura=24)
            r += 1
        elif tipo in ("TEXTO",):
            U.mesclar(ws, f"B{r}:F{r+1}", a, "NOTA_WRAP")
            r += 2
        elif tipo == "TABELA":
            U.mesclar(ws, f"B{r}:C{r}", a, "CALC_TXT")
            U.mesclar(ws, f"D{r}:F{r}", b, "NOTA")
            r += 1
        elif tipo == "COR":
            U.mesclar(ws, f"B{r}:C{r}", a, estilo or "CALC_TXT_C")
            U.mesclar(ws, f"D{r}:F{r}", b, "NOTA")
            r += 1
        elif tipo == "PASSO":
            U.mesclar(ws, f"B{r}:C{r}", a, "KPI_LABEL")
            U.mesclar(ws, f"D{r}:F{r}", b, "NOTA")
            r += 1
        elif tipo == "FAQ":
            U.mesclar(ws, f"B{r}:C{r}", a, "CALC_TXT")
            U.mesclar(ws, f"D{r}:F{r}", b, "NOTA")
            r += 1
    U.mesclar(ws, f"B{r+1}:F{r+2}",
              "Dúvidas de uso: comece sempre pela aba LANÇAMENTOS — todas as demais telas (dashboard, "
              "DRE, fluxo, orçamento, indicadores e relatórios) são calculadas a partir dela.", "NOTA_WRAP")
    U.preparar(ws, grade=False, tab_cor=E.GRAY_TEXT, area_impressao=f"A1:F{r+3}")
    ctx.instrucoes = ws
    return ws
