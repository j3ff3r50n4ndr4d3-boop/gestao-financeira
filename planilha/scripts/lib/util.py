"""
util.py — utilitários de construção das abas (células, seções, tabelas,
validação de dados, formatação condicional, intervalos nomeados).
"""
from openpyxl.formatting.rule import (CellIsRule, ColorScaleRule, DataBarRule,
                                      FormulaRule, IconSetRule)
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import estilos as E


# ------------------------------------------------------------------ refs
def L(col: int) -> str:
    """1 -> A, 27 -> AA"""
    return get_column_letter(col)


def R(col: int, row: int, abs_col: bool = False, abs_row: bool = False) -> str:
    c = L(col)
    if abs_col:
        c = "$" + c
    r = str(row)
    if abs_row:
        r = "$" + r
    return f"{c}{r}"


def faixa(col1, row1, col2, row2, aba=None):
    """Intervalo A1:B2 (opcionalmente com aba)."""
    ref = f"{R(col1, row1)}:{R(col2, row2)}"
    if aba:
        return f"'{aba}'!{ref}" if (" " in aba or "-" in aba or "Ç" in aba) else f"{aba}!{ref}"
    return ref


def ref_aba(aba, celula):
    """'DRE GERENCIAL'!$C$10"""
    if any(ch in aba for ch in " -+'*%$"):
        return f"'{aba}'!{celula}"
    return f"{aba}!{celula}"


def nome(wb, rotulo: str, ref: str):
    """Cria um intervalo nomeado global."""
    from openpyxl.workbook.defined_name import DefinedName
    dn = DefinedName(rotulo, attr_text=ref)
    if hasattr(wb.defined_names, "add"):
        try:
            wb.defined_names.add(dn)
        except AttributeError:
            wb.defined_names[rotulo] = dn
    else:  # openpyxl < 3.1
        wb.defined_names.append(dn)
    return rotulo


# ------------------------------------------------------------------ escrita
def put(ws, row, col, value=None, style=None, fmt=None, com_valor=True):
    c = ws.cell(row=row, column=col)
    if com_valor and value is not None:
        c.value = value
    if style:
        c.style = style
    if fmt:
        c.number_format = fmt
    return c


def put_ref(ws, ref, value=None, style=None, fmt=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if style:
        c.style = style
    if fmt:
        c.number_format = fmt
    return c


def estilizar(ws, rng, style):
    for row in ws[rng]:
        for c in row:
            c.style = style


def mesclar(ws, rng, value=None, style=None, fmt=None):
    ws.merge_cells(rng)
    ini = rng.split(":")[0]
    c = ws[ini]
    if value is not None:
        c.value = value
    if style:
        c.style = style
    if fmt:
        c.number_format = fmt
    # garante estilo na área mesclada
    if style:
        estilizar(ws, rng, style)
    return c


def preencher(ws, col, row_ini, row_fim, template, style):
    """Escreve a mesma fórmula (com {r}) em uma coluna."""
    for r in range(row_ini, row_fim + 1):
        ws.cell(row=r, column=col, value=template.format(r=r)).style = style


def preencher_vals(ws, col, row_ini, row_fim, valores, style):
    for i, r in enumerate(range(row_ini, row_fim + 1)):
        v = valores[i] if i < len(valores) else None
        ws.cell(row=r, column=col, value=v).style = style


def vazio_estilo(ws, col, row_ini, row_fim, style):
    preencher_vals(ws, col, row_ini, row_fim, [], style)


def titulo(ws, row, texto, subtitulo=None, span=8, inicio=1, cor=None):
    c = put(ws, row, inicio, texto, "TITULO")
    if span > 1:
        ws.merge_cells(start_row=row, start_column=inicio, end_row=row, end_column=inicio + span - 1)
    ws.row_dimensions[row].height = 24
    if subtitulo:
        c2 = put(ws, row + 1, inicio, subtitulo, "SUBTITULO")
        ws.merge_cells(start_row=row + 1, start_column=inicio, end_row=row + 1, end_column=inicio + span - 1)
    return c


def secao(ws, row, texto, span=8, col=1, claro=False, altura=20):
    est = "SECAO_CLARA" if claro else "SECAO"
    if span > 1:
        mesclar(ws, f"{R(col,row)}:{R(col+span-1,row)}", texto, est)
    else:
        put(ws, row, col, texto, est)
    ws.row_dimensions[row].height = altura


def cabecalho(ws, row, col_ini, headers, style="TH", altura=32):
    for i, h in enumerate(headers):
        put(ws, row, col_ini + i, h, style)
    ws.row_dimensions[row].height = altura


def cartao(ws, row, col, span, label, valor, style_valor="KPI_VALUE", nota=None,
           altura_label=26, altura_valor=28):
    """Cartão de KPI: linha de rótulo, linha de valor e (opcional) linha de nota."""
    mesclar(ws, f"{R(col,row)}:{R(col+span-1,row)}", label, "KPI_LABEL")
    mesclar(ws, f"{R(col,row+1)}:{R(col+span-1,row+1)}", valor, style_valor)
    ws.row_dimensions[row].height = altura_label
    ws.row_dimensions[row + 1].height = altura_valor
    if nota:
        mesclar(ws, f"{R(col,row+2)}:{R(col+span-1,row+2)}", nota, "KPI_NOTE")
        ws.row_dimensions[row + 2].height = 15


# ------------------------------------------------------------------ validação
def dv_lista(ws, rng, formula, titulo="Valor inválido", msg="Selecione um valor da lista.",
             estilo="stop", permite_vazio=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=permite_vazio,
                        showErrorMessage=True, showInputMessage=True,
                        errorTitle=titulo, error=msg, errorStyle=estilo,
                        promptTitle="Valores permitidos", prompt=msg)
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def dv_decimal(ws, rng, op="greaterThanOrEqual", f1="0", titulo="Valor inválido",
               msg="Informe um número válido (>= 0).", permite_vazio=True, estilo="stop"):
    dv = DataValidation(type="decimal", operator=op, formula1=f1, allow_blank=permite_vazio,
                        showErrorMessage=True, errorTitle=titulo, error=msg, errorStyle=estilo)
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def dv_data(ws, rng, op="between", f1="DATE(2000,1,1)", f2="DATE(2100,12,31)",
            titulo="Data inválida", msg="Informe uma data válida (entre 01/01/2000 e 31/12/2100)."):
    dv = DataValidation(type="date", operator=op, formula1=f1, formula2=f2, allow_blank=True,
                        showErrorMessage=True, errorTitle=titulo, error=msg, errorStyle="warning")
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def dv_inteiro(ws, rng, op="between", f1="0", f2="999", titulo="Número inválido",
               msg="Informe um número inteiro válido."):
    dv = DataValidation(type="whole", operator=op, formula1=f1, formula2=f2, allow_blank=True,
                        showErrorMessage=True, errorTitle=titulo, error=msg, errorStyle="stop")
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


# ------------------------------------------------------------------ formatação condicional
def cf_texto(ws, rng, texto, bg, tx=None, op="equal"):
    f = [f'"{texto}"']
    r = CellIsRule(operator=op, formula=f,
                   fill=PatternFill("solid", fgColor=bg),
                   font=Font(color=tx) if tx else None)
    ws.conditional_formatting.add(rng, r)


def cf_contem(ws, rng, texto, bg, tx=None):
    r = FormulaRule(formula=[f'ISNUMBER(SEARCH("{texto}",{rng.split(":")[0]}))'],
                    fill=PatternFill("solid", fgColor=bg),
                    font=Font(color=tx) if tx else None, stopIfTrue=False)
    ws.conditional_formatting.add(rng, r)


def cf_maior(ws, rng, valor, bg, tx=None, op="greaterThan"):
    r = CellIsRule(operator=op, formula=[str(valor)],
                   fill=PatternFill("solid", fgColor=bg) if bg else None,
                   font=Font(color=tx, bold=True if tx else None) if tx else None)
    ws.conditional_formatting.add(rng, r)


def cf_formula(ws, rng, formula, bg=None, tx=None, bold=None):
    r = FormulaRule(formula=[formula],
                    fill=PatternFill("solid", fgColor=bg) if bg else None,
                    font=Font(color=tx, bold=bold) if (tx or bold) else None)
    ws.conditional_formatting.add(rng, r)


def cf_barras(ws, rng, cor=E.BLUE):
    ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0,
                                                   end_type="max", color=cor, showValue=True))


def cf_escala(ws, rng, c1=E.RED_BG, c2=E.YEL_BG, c3=E.GREEN_BG):
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=c1,
                                                      mid_type="percentile", mid_value=50, mid_color=c2,
                                                      end_type="max", end_color=c3))


def cf_sem_erro(ws, rng):  # destaca textos de erro/validação
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("⚠",{rng.split(":")[0]}))'],
        fill=PatternFill("solid", fgColor=E.RED_BG),
        font=Font(color=E.RED_TX, bold=True)))


def larguras(ws, mapa):
    """Define larguras de coluna: {'A': 12, 'B': 30}"""
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


def alturas(ws, mapa):
    for row, h in mapa.items():
        ws.row_dimensions[row].height = h


# ------------------------------------------------------------------ setup de aba
def preparar(ws, freeze=None, filtro=None, grade=False, zoom=100, tab_cor=None,
             paisagem=True, area_impressao=None, repetir_linhas=None, ajustar_largura=1,
             rodape=None):
    if freeze:
        ws.freeze_panes = freeze
    if filtro:
        ws.auto_filter.ref = filtro
    ws.sheet_view.showGridLines = grade
    ws.sheet_view.zoomScale = zoom
    if tab_cor:
        ws.sheet_properties.tabColor = tab_cor
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.page_setup.fitToWidth = ajustar_largura
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if repetir_linhas:
        ws.print_title_rows = repetir_linhas
    if area_impressao:
        ws.print_area = area_impressao
    ws.oddFooter.left.text = rodape or "&A"
    ws.oddFooter.right.text = "Página &P de &N"
