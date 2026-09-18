"""
estilos.py — Paleta, formatos numéricos e estilos nomeados da planilha.

Regras de design:
  * Layout minimalista, hierarquia visual clara, cores usadas apenas para
    indicar status / alertas (vermelho, amarelo, verde) e para a identidade
    visual da planilha (azul-marinho / azul).
  * Células AMARELAS  = entrada de dados do usuário (desbloqueadas)
  * Células CINZA     = fórmulas automáticas (bloqueadas)
  * Formatos numéricos em padrão en-US ("#,##0.00"), que o Excel e o Google
    Sheets exibem automaticamente no padrão brasileiro (R$ 1.234,56) conforme
    o idioma/localidade do arquivo.
"""
from openpyxl.styles import (Alignment, Border, Font, NamedStyle, PatternFill,
                             Protection, Side)

# --------------------------------------------------------------------------
# PALETA
# --------------------------------------------------------------------------
NAVY        = "1F3864"   # azul-marinho (títulos, cabeçalhos)
NAVY_DARK   = "13233F"
BLUE        = "2E75B6"   # azul médio (destaques, seções)
BLUE_LIGHT  = "D9E2F3"   # azul claro (faixas de seção)
BLUE_PALE   = "EEF3FA"   # azul muito claro (linhas alternadas)
GRAY_BG     = "F2F2F2"   # cinza cálculo
WHITE       = "FFFFFF"
GRAY_LINE   = "BFBFBF"
GRAY_TEXT   = "595959"
AMBER_BG    = "FFF2CC"   # amarelo = célula de entrada
AMBER_LINE  = "D6B656"
GREEN_BG    = "C6EFCE"
GREEN_TX    = "0B6623"
RED_BG      = "FFC7CE"
RED_TX      = "9C0006"
YEL_BG      = "FFEB9C"
YEL_TX      = "9C6500"
NEUTRAL_BG  = "E7E6E6"
TEAL        = "2E8B8B"
ORANGE      = "ED7D31"

# --------------------------------------------------------------------------
# FORMATOS NUMÉRICOS
# --------------------------------------------------------------------------
FMT_MONEY   = 'R$ #,##0.00'
FMT_MONEY0  = 'R$ #,##0'
FMT_PCT     = '0.0%'
FMT_PCT0    = '0%'
FMT_DATE    = 'DD/MM/YYYY'
FMT_DATE_S  = 'DD/MM'
FMT_INT     = '#,##0'
FMT_INT2    = '#,##0.00'
FMT_TXT     = '@'

# --------------------------------------------------------------------------
# BORDAS / ALINHAMENTOS
# --------------------------------------------------------------------------
_borda_fina   = Side(style="thin", color=GRAY_LINE)
_borda_ambar  = Side(style="thin", color=AMBER_LINE)
_borda_baixa  = Side(style="thin", color="E2E2E2")

BORDA = Border(left=_borda_fina, right=_borda_fina, top=_borda_fina, bottom=_borda_fina)
BORDA_INPUT = Border(left=_borda_ambar, right=_borda_ambar, top=_borda_ambar, bottom=_borda_ambar)
BORDA_LEVE = Border(left=_borda_baixa, right=_borda_baixa, top=_borda_baixa, bottom=_borda_baixa)
BORDA_TOPO = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="double", color=NAVY))

LOCK   = Protection(locked=True)
UNLOCK = Protection(locked=False)

AL_CENTER = Alignment(horizontal="center", vertical="center")
AL_LEFT   = Alignment(horizontal="left", vertical="center")
AL_RIGHT  = Alignment(horizontal="right", vertical="center")
AL_WRAP   = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)

# --------------------------------------------------------------------------
# DEFINIÇÃO DOS ESTILOS NOMEADOS
# --------------------------------------------------------------------------
# (nome, fonte, preenchimento, borda, alinhamento, formato, proteção)
DEFS = [
    # ---------- títulos e seções ----------
    ("TITULO",       Font(name="Calibri", size=18, bold=True, color=NAVY),      None, None, AL_LEFT, FMT_TXT, LOCK),
    ("SUBTITULO",    Font(name="Calibri", size=10, color=GRAY_TEXT),            None, None, AL_LEFT, FMT_TXT, LOCK),
    ("SECAO",        Font(name="Calibri", size=11, bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY), None, AL_LEFT, FMT_TXT, LOCK),
    ("SECAO_CLARA",  Font(name="Calibri", size=11, bold=True, color=NAVY),  PatternFill("solid", fgColor=BLUE_LIGHT), None, AL_LEFT, FMT_TXT, LOCK),
    ("BLOCO",        Font(name="Calibri", size=10, bold=True, color=NAVY),  PatternFill("solid", fgColor=BLUE_PALE), BORDA, AL_LEFT, FMT_TXT, LOCK),
    # ---------- cabeçalhos de tabela ----------
    ("TH",           Font(name="Calibri", size=10, bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY), BORDA, AL_CENTER_WRAP, FMT_TXT, LOCK),
    ("TH2",          Font(name="Calibri", size=10, bold=True, color=NAVY),  PatternFill("solid", fgColor=BLUE_LIGHT), BORDA, AL_CENTER_WRAP, FMT_TXT, LOCK),
    # ---------- entradas do usuário (amarelo = digite aqui) ----------
    ("IN_TXT",       Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_LEFT,   FMT_TXT,     UNLOCK),
    ("IN_TXT_C",     Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_CENTER, FMT_TXT,     UNLOCK),
    ("IN_MONEY",     Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_RIGHT,  FMT_MONEY,   UNLOCK),
    ("IN_MONEY0",    Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_RIGHT,  FMT_MONEY0,  UNLOCK),
    ("IN_PCT",       Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_CENTER, FMT_PCT,     UNLOCK),
    ("IN_INT",       Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_CENTER, FMT_INT,     UNLOCK),
    ("IN_DATE",      Font(name="Calibri", size=10, color="00305B"),   PatternFill("solid", fgColor=AMBER_BG), BORDA_INPUT, AL_CENTER, FMT_DATE,    UNLOCK),
    # ---------- campos calculados (cinza = não editar) ----------
    ("CALC_TXT",     Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_LEFT,   FMT_TXT,   LOCK),
    ("CALC_TXT_C",   Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_CENTER, FMT_TXT,   LOCK),
    ("CALC_MONEY",   Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_RIGHT,  FMT_MONEY, LOCK),
    ("CALC_MONEY0",  Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_RIGHT,  FMT_MONEY0, LOCK),
    ("CALC_PCT",     Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_CENTER, FMT_PCT,   LOCK),
    ("CALC_INT",     Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_CENTER, FMT_INT,   LOCK),
    ("CALC_DATE",    Font(name="Calibri", size=10, color="404040"),   PatternFill("solid", fgColor=GRAY_BG),  BORDA, AL_CENTER, FMT_DATE,  LOCK),
    ("CALC_TXT_SB",  Font(name="Calibri", size=10, color="404040"),   None, BORDA, AL_LEFT, FMT_TXT, LOCK),
    # ---------- linhas de total ----------
    ("TOTAL_TXT",    Font(name="Calibri", size=10, bold=True, color=NAVY), PatternFill("solid", fgColor=BLUE_LIGHT), BORDA, AL_LEFT,  FMT_TXT,     LOCK),
    ("TOTAL_MONEY",  Font(name="Calibri", size=10, bold=True, color=NAVY), PatternFill("solid", fgColor=BLUE_LIGHT), BORDA, AL_RIGHT, FMT_MONEY,   LOCK),
    ("TOTAL_PCT",    Font(name="Calibri", size=10, bold=True, color=NAVY), PatternFill("solid", fgColor=BLUE_LIGHT), BORDA, AL_CENTER,FMT_PCT,     LOCK),
    ("TOTAL_INT",    Font(name="Calibri", size=10, bold=True, color=NAVY), PatternFill("solid", fgColor=BLUE_LIGHT), BORDA, AL_CENTER,FMT_INT,     LOCK),
    # ---------- cartões de KPI ----------
    ("KPI_LABEL",    Font(name="Calibri", size=9,  bold=True, color=GRAY_TEXT), PatternFill("solid", fgColor=BLUE_PALE), Border(left=_borda_fina, right=_borda_fina, top=_borda_fina), AL_CENTER_WRAP, FMT_TXT,   LOCK),
    ("KPI_VALUE",    Font(name="Calibri", size=14, bold=True, color=NAVY),      PatternFill("solid", fgColor=BLUE_PALE), Border(left=_borda_fina, right=_borda_fina),                  AL_CENTER,      FMT_MONEY, LOCK),
    ("KPI_VALUE0",   Font(name="Calibri", size=14, bold=True, color=NAVY),      PatternFill("solid", fgColor=BLUE_PALE), Border(left=_borda_fina, right=_borda_fina),                  AL_CENTER,      FMT_MONEY0,LOCK),
    ("KPI_PCT",      Font(name="Calibri", size=14, bold=True, color=NAVY),      PatternFill("solid", fgColor=BLUE_PALE), Border(left=_borda_fina, right=_borda_fina),                  AL_CENTER,      FMT_PCT,    LOCK),
    ("KPI_VALUE_TXT",Font(name="Calibri", size=11, bold=True, color=NAVY),      PatternFill("solid", fgColor=BLUE_PALE), Border(left=_borda_fina, right=_borda_fina),                  AL_CENTER,      FMT_TXT,    LOCK),
    ("KPI_NOTE",     Font(name="Calibri", size=8,  color=GRAY_TEXT),            PatternFill("solid", fgColor=BLUE_PALE), Border(left=_borda_fina, right=_borda_fina, bottom=_borda_fina), AL_CENTER_WRAP, FMT_TXT,   LOCK),
    # ---------- textos auxiliares ----------
    ("NOTA",         Font(name="Calibri", size=9, italic=True, color=GRAY_TEXT), None, None, AL_LEFT, FMT_TXT, LOCK),
    ("NOTA_WRAP",    Font(name="Calibri", size=9, color=GRAY_TEXT), None, None, AL_WRAP, FMT_TXT, LOCK),
    ("TXT",          Font(name="Calibri", size=10), None, None, AL_LEFT,   FMT_TXT,   LOCK),
    ("TXT_C",        Font(name="Calibri", size=10), None, None, AL_CENTER, FMT_TXT,   LOCK),
    ("TXT_R",        Font(name="Calibri", size=10), None, None, AL_RIGHT,  FMT_TXT,   LOCK),
    ("TXT_B",        Font(name="Calibri", size=10, bold=True, color=NAVY), None, None, AL_LEFT, FMT_TXT, LOCK),
    ("LINK",         Font(name="Calibri", size=10, bold=True, color=BLUE, underline="single"), None, BORDA, AL_CENTER, FMT_TXT, LOCK),
    ("MENU_BOTAO",   Font(name="Calibri", size=11, bold=True, color=WHITE), PatternFill("solid", fgColor=BLUE), BORDA, AL_CENTER_WRAP, FMT_TXT, LOCK),
    ("MENU_BOTAO_D", Font(name="Calibri", size=11, bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY), BORDA, AL_CENTER_WRAP, FMT_TXT, LOCK),
    # ---------- números/percentuais simples (sem borda) ----------
    ("N_MONEY",      Font(name="Calibri", size=10), None, None, AL_RIGHT,  FMT_MONEY, LOCK),
    ("N_MONEY0",     Font(name="Calibri", size=10), None, None, AL_RIGHT,  FMT_MONEY0,LOCK),
    ("N_PCT",        Font(name="Calibri", size=10), None, None, AL_CENTER, FMT_PCT,   LOCK),
    ("N_INT",        Font(name="Calibri", size=10), None, None, AL_CENTER, FMT_INT,   LOCK),
    ("N_DATE",       Font(name="Calibri", size=10), None, None, AL_CENTER, FMT_DATE,  LOCK),
    ("N_DATE_S",     Font(name="Calibri", size=10), None, None, AL_CENTER, FMT_DATE_S,LOCK),
    # ---------- dashboard / painel ----------
    ("TITULO_DASH",   Font(name="Calibri", size=18, bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY), None, AL_CENTER_WRAP, FMT_TXT, LOCK),
    ("SUBTITULO_DASH",Font(name="Calibri", size=10, italic=True, color=WHITE), PatternFill("solid", fgColor=BLUE), None, AL_CENTER_WRAP, FMT_TXT, LOCK),
    ("FILTRO_LABEL",  Font(name="Calibri", size=9, bold=True, color=NAVY), PatternFill("solid", fgColor=BLUE_LIGHT), BORDA, AL_CENTER_WRAP, FMT_TXT, LOCK),
    ("ALERTA_TXT",    Font(name="Calibri", size=10, color=GRAY_TEXT), PatternFill("solid", fgColor=BLUE_PALE), BORDA, AL_LEFT, FMT_TXT, LOCK),
    ("ALERTA_ICONE",  Font(name="Calibri", size=10, bold=True, color=NAVY), PatternFill("solid", fgColor=BLUE_PALE), BORDA, AL_CENTER, FMT_TXT, LOCK),
    ("LINK_ABA",      Font(name="Calibri", size=10, bold=True, color=BLUE, underline="single"), None, BORDA, AL_CENTER, FMT_TXT, LOCK),
]


def registrar_estilos(wb):
    """Registra todos os estilos nomeados no workbook."""
    for nome, fonte, fill, borda, alin, fmt, prot in DEFS:
        ns = NamedStyle(name=nome)
        ns.font = fonte
        if fill is not None:
            ns.fill = fill
        if borda is not None:
            ns.border = borda
        ns.alignment = alin
        ns.number_format = fmt
        ns.protection = prot
        wb.add_named_style(ns)
    return wb


def larguras(ws, mapa):
    """Define larguras de coluna. mapa = {'A': 12, 'B': 20, ...}"""
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


def alturas(ws, mapa):
    for row, h in mapa.items():
        ws.row_dimensions[row].height = h
